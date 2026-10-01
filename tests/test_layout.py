"""Layout behavior with synthetic images and text; actual page QA is separate."""
import contextlib
import io
from pathlib import Path
import sys
import tempfile
import unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'snu-report-core/scripts'))
import build
from docx_post import (frame_all, tidy_cells, keep_related_content, force_black_text,
                       clean_paragraphs, para_center_tight, fit_images)
from make_template import build_reference

class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.style = yaml.safe_load((ROOT / 'snu-report-core/templates/style.yaml').read_text(encoding='utf-8'))

    def tearDown(self):
        self.temp.cleanup()

    def test_missing_photo_is_editable_text_with_resolved_reference(self):
        with contextlib.redirect_stderr(io.StringIO()):
            raw, missing = build.photo_placeholders('See @fig:a.\n\n![Missing photo](figs/a.png){#fig:a hint="Photo instruction"}', self.folder, self.folder, self.style)
        rendered, warnings = build.preprocess(raw, self.style)
        self.assertEqual(missing, ['figs/a.png'])
        self.assertFalse(warnings)
        self.assertIn('See Fig.1.', rendered)
        self.assertIn('미제공: Fig.1 - Missing photo', rendered)
        self.assertNotIn('![', rendered)
        self.assertNotIn('Photo instruction', rendered)
        self.assertFalse((self.folder / 'build/placeholders').exists())

    def test_portrait_frame_does_not_expand_to_page_width(self):
        image = self.folder / 'portrait.png'
        Image.new('RGB', (200,400), 'white').save(image)
        doc = Document()
        picture = doc.add_paragraph(style='Normal')
        picture.style = doc.styles.add_style('Figure', 1)
        picture.add_run().add_picture(str(image), width=Mm(50))
        frame_all(doc, self.style, 9000)
        width = int(doc.tables[0]._tbl.find('.//' + qn('w:tblW')).get(qn('w:w')))
        self.assertLess(width, 3500)
        self.assertGreater(width, 2800)

    def test_long_table_can_continue_with_repeated_header_and_padding(self):
        doc = Document()
        table = doc.add_table(rows=12, cols=2)
        tidy_cells(table._tbl, 10.5)
        self.assertIsNotNone(table.rows[0]._tr.find('.//' + qn('w:tblHeader')))
        for row in table.rows:
            for cell in row.cells:
                self.assertFalse(cell.paragraphs[0].paragraph_format.keep_with_next)
                self.assertEqual(cell._tc.find('.//' + qn('w:tcMar') + '/' + qn('w:top')).get(qn('w:w')), '100')

    def test_equation_stays_with_its_introduction(self):
        doc = Document()
        intro = doc.add_paragraph('The transfer function follows.')
        equation = doc.add_paragraph()
        equation._p.append(OxmlElement('m:oMathPara'))
        last = doc.add_paragraph('A separate interpretation.')
        self.assertEqual(keep_related_content(doc), 1)
        self.assertTrue(intro.paragraph_format.keep_with_next)
        self.assertFalse(last.paragraph_format.keep_with_next)

    def test_short_result_stays_with_calculation_and_its_introduction(self):
        doc = Document()
        result = doc.add_paragraph('Measured delay: 20 us.')
        intro = doc.add_paragraph('The phase follows from the measured delay.')
        equation = doc.add_paragraph()
        equation._p.append(OxmlElement('m:oMathPara'))
        self.assertEqual(keep_related_content(doc), 2)
        self.assertTrue(result.paragraph_format.keep_with_next)
        self.assertTrue(intro.paragraph_format.keep_with_next)

    def test_two_short_setup_paragraphs_stay_with_the_photo(self):
        image = self.folder / 'board.png'
        Image.new('RGB', (100,200), 'white').save(image)
        doc = Document()
        setup = doc.add_paragraph('Circuit setup.')
        conditions = doc.add_paragraph('Input and supply conditions.')
        table = doc.add_table(rows=2, cols=1)
        table.cell(0,0).paragraphs[0].add_run().add_picture(str(image), width=Mm(40))
        table.cell(1,0).text = 'Fig.1 - board'
        self.assertEqual(keep_related_content(doc), 2)
        self.assertTrue(setup.paragraph_format.keep_with_next)
        self.assertTrue(conditions.paragraph_format.keep_with_next)

    def test_missing_photo_notice_does_not_orphan_after_a_table(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        caption = table.cell(1,0).paragraphs[0]
        caption.text = 'Table.1 - theory'
        style = doc.styles.add_style('Missing Figure', 1)
        doc.add_paragraph('Missing measured photo.', style=style)
        self.assertEqual(keep_related_content(doc), 1)
        self.assertTrue(caption.paragraph_format.keep_with_next)

    def test_page_field_uses_simple_field(self):
        path = build_reference(ROOT / 'snu-report-core/templates/style.yaml', self.folder / 'reference.docx', {})
        doc = Document(path)
        fields = doc.sections[0].footer._element.findall('.//' + qn('w:fldSimple'))
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0].get(qn('w:instr')).strip(), 'PAGE')

    def test_caption_explicitly_disables_inherited_keep_next(self):
        doc = Document()
        doc.styles['Caption'].paragraph_format.keep_with_next = True
        caption = doc.add_paragraph('Fig.1 - example', style='Caption')
        para_center_tight(caption._p, keep_next=False)
        self.assertFalse(caption.paragraph_format.keep_with_next)

    def test_black_text_removes_theme_and_link_colors_in_all_text_parts(self):
        from docx.enum.dml import MSO_THEME_COLOR_INDEX
        doc = Document()
        body = doc.add_paragraph('Body').runs[0]
        body.font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_1
        try:
            hyperlink = doc.styles['Hyperlink']
        except KeyError:
            hyperlink = doc.styles.add_style('Hyperlink', 2)
        hyperlink.font.color.theme_color = MSO_THEME_COLOR_INDEX.HYPERLINK
        doc.sections[0].header.paragraphs[0].add_run('Header').font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_2
        doc.sections[0].footer.paragraphs[0].add_run('Footer').font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_3
        force_black_text(doc)
        for root in (doc.element, doc.styles.element, doc.sections[0].header._element, doc.sections[0].footer._element):
            for color in root.iter(qn('w:color')):
                self.assertEqual(dict(color.attrib), {qn('w:val'):'000000'})

    def test_empty_enters_and_body_line_breaks_are_normalized(self):
        doc = Document()
        doc.add_paragraph('')
        doc.add_paragraph('')
        paragraph = doc.add_paragraph('First line')
        paragraph.add_run().add_break()
        paragraph.add_run('Second line')
        clean_paragraphs(doc, self.style)
        self.assertEqual(len(doc.paragraphs), 1)
        self.assertEqual(paragraph.text, 'First line Second line')
        self.assertEqual(paragraph.paragraph_format.space_after.pt, 6)
        self.assertIsNone(paragraph._p.find('.//' + qn('w:br')))

    def test_absolute_photo_caps_override_requested_full_width(self):
        image = self.folder / 'large.png'
        Image.new('RGB', (1000,2000), 'white').save(image)
        doc = Document()
        photo = doc.add_paragraph().add_run().add_picture(str(image), width=Mm(200))
        fit_images(doc, int(Mm(140)), int(Mm(95)), int(Mm(80)))
        self.assertLessEqual(photo.width, Mm(140))
        self.assertLessEqual(photo.height, Mm(80))

if __name__ == '__main__':
    unittest.main()
