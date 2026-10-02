"""Synthetic fixtures only. No student reports, photos or course PDFs."""
import argparse
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "snu-report-core/scripts"))
import bootstrap
import build
import quality
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image
import yaml


class RuntimeTests(unittest.TestCase):
    def options(self, agent, skills_only=False):
        return argparse.Namespace(agent=agent, skills_only=skills_only, no_ltspice=False,
                                  deps_only=False, preview=False)

    def test_host_session_detection(self):
        self.assertEqual(bootstrap.selected_agent("auto", {"CODEX_THREAD_ID": "x"}), "codex")
        self.assertEqual(bootstrap.selected_agent("auto", {"CLAUDECODE": "1"}), "claude")
        for env in ({}, {"CODEX_HOME": "/installed/codex"}, {"CODEX_THREAD_ID": "x", "CLAUDECODE": "1"}):
            with self.assertRaises(ValueError):
                bootstrap.selected_agent("auto", env)
        self.assertEqual(bootstrap.selected_agent("codex", {"CLAUDECODE": "1"}), "codex")

    @patch.object(bootstrap, "claude_code", return_value="claude")
    @patch.object(bootstrap, "claude_desktop_config", return_value=Path("claude.json"))
    @patch.object(bootstrap, "WIN", True)
    def test_codex_never_targets_claude(self, *_):
        applicable = [name for name, _, applies, *_ in bootstrap.steps(self.options("codex")) if applies]
        self.assertTrue(any("Codex" in name for name in applicable))
        self.assertFalse(any("Claude" in name for name in applicable))

    @patch.object(bootstrap, "claude_code", return_value="claude")
    @patch.object(bootstrap, "WIN", True)
    def test_claude_never_targets_codex(self, *_):
        applicable = [name for name, _, applies, *_ in bootstrap.steps(self.options("claude")) if applies]
        self.assertFalse(any("Codex" in name for name in applicable))

    def test_skills_only_has_no_dependencies_or_mcp(self):
        self.assertTrue(all(name.startswith("스킬:") for name, *_ in bootstrap.steps(self.options("codex", True))))

    def test_installer_updates_files_and_preserves_user_files(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            source, target = base / "source", base / "target"
            (source / "snu-test/references").mkdir(parents=True)
            (source / "snu-test/SKILL.md").write_text("version 1")
            (source / "snu-test/references/runtime.md").write_text("codex")
            with patch.object(bootstrap, "SKILLS_DIR", source), patch.object(bootstrap, "SKILL_NAMES", ["snu-test"]), patch.object(bootstrap, "SKILL_ROOT", target), patch.object(bootstrap, "DRY", False):
                self.assertFalse(bootstrap.codex_skills_installed())
                bootstrap.do_codex_skills()
                self.assertTrue(bootstrap.codex_skills_installed())
                (target / "snu-test/user-notes.txt").write_text("keep")
                (source / "snu-test/references/runtime.md").write_text("updated")
                self.assertFalse(bootstrap.codex_skills_installed())
                bootstrap.do_codex_skills()
                self.assertEqual((target / "snu-test/user-notes.txt").read_text(), "keep")
                self.assertEqual((target / "snu-test/references/runtime.md").read_text(), "updated")

    def test_real_discovery_folders_and_links(self):
        for host in (".agents", ".claude"):
            for slot in (ROOT / host / "skills").iterdir():
                self.assertTrue(slot.is_dir())
                entry = slot / "SKILL.md"
                front = yaml.safe_load(entry.read_text(encoding="utf-8").split("---", 2)[1])
                self.assertEqual(front["name"], slot.name)
                host_name = 'codex' if host == '.agents' else 'claude'
                other_host = 'claude' if host == '.agents' else 'codex'
                text = entry.read_text(encoding='utf-8')
                if slot.name == 'snu-lab-photo':
                    self.assertIn(host_name, text.lower())
                    self.assertNotIn(other_host, text.lower())
                else:
                    self.assertIn('runtime-' + host_name + '.md', text)
                    self.assertNotIn('runtime-' + other_host + '.md', text)
                for destination in __import__('re').findall(r"\]\(([^)]+)\)", entry.read_text(encoding="utf-8")):
                    self.assertTrue((entry.parent / destination).is_file(), destination)


class QualityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="snu-fixture-")
        self.ws = Path(self.temp.name)
        self.lab = self.ws / "courses/circuit/lab01"
        (self.lab / "report").mkdir(parents=True)
        (self.ws / "profile.yaml").write_text('name: Test\nstudent_id: 2099-00000\n', encoding="utf-8")
        (self.lab / "meta.yaml").write_text('lab: 1\ntitle: Synthetic RC exercise\n', encoding="utf-8")
        (self.lab / "source.txt").write_text("Synthetic assignment: calculate RC gain.", encoding="utf-8")
        self.raw = "# (1) RC gain\n\nThe calculated gain is 0.5.\n"
        (self.lab / "report/report.md").write_text(self.raw, encoding="utf-8")
        self.requirements = {"scope": {"report": {"status":"confirmed", "source":"source.txt", "locator":"line 1"}},
            "items": [{"id":"(1)", "kind":"report", "status":"answered", "source":"source.txt", "locator":"line 1", "answer":"RC gain", "evidence":["source.txt"]}]}
        self.write_requirements()
        self.docx = self.lab / "build/lab01_2099-00000_Test.docx"
        self.docx.parent.mkdir()
        self.document = Document()
        self.document.add_heading("(1) RC gain", 1)
        self.document.add_paragraph("The calculated gain is 0.5.")
        self.document.save(self.docx)

    def tearDown(self):
        self.temp.cleanup()

    def write_requirements(self):
        (self.lab / "requirements.yaml").write_text(yaml.safe_dump(self.requirements), encoding="utf-8")

    def create_render_fixture(self, count=2):
        # These images are dummy renderer output; this tests the gate, not layout.
        folder = quality.qa_dir(self.docx)
        folder.mkdir(parents=True, exist_ok=True)
        pages = {}
        for i in range(1, count+1):
            path = folder / f"page-{i}.png"
            Image.new("RGB", (80, 100), "white").save(path)
            pages[str(i)] = quality.digest(path)
        quality.save_json(folder / "render.json", {"docx_sha256":quality.digest(self.docx), "pages":pages, "pdf_sha256":None})

    def record_build(self):
        quality.record_build(self.docx, self.lab, "report", ROOT / "snu-report-core/templates/style.yaml")

    def test_good_requirements(self):
        self.assertEqual(quality.check_requirements(self.lab, "report", self.raw), [])

    def test_duplicate_heading_number_blocks_delivery(self):
        from proof import deduplicate_heading_labels
        heading = self.document.paragraphs[0]
        for text in ('2.3) 3) 120 nF', '1. (1) Circuit', '1.1) 가) Transfer'):
            heading.text = text
            self.document.save(self.docx)
            self.assertTrue(any('제목 번호 중복' in e for e in quality.check_docx(self.docx)['errors']))
        for text in ('2.3) 120 nF', '2.2) 47 nF', '1. 3-bit circuit'):
            heading.text = text
            self.document.save(self.docx)
            self.assertFalse(quality.check_docx(self.docx)['errors'])
            self.assertEqual(deduplicate_heading_labels(text), text)
        self.assertEqual(deduplicate_heading_labels('2.3) 3) 120 nF'), '2.3) 120 nF')

    def test_unconfirmed_scope_and_missing_source_block(self):
        self.requirements["scope"]["report"]["status"] = "provisional"
        self.requirements["items"][0]["source"] = "missing.txt"
        self.write_requirements()
        self.assertGreaterEqual(len(quality.check_requirements(self.lab, "report", self.raw)), 2)

    def test_answer_and_evidence_are_verified(self):
        self.requirements["items"][0].update(answer="absent heading", evidence=["absent.csv"])
        self.write_requirements()
        self.assertEqual(len(quality.check_requirements(self.lab, "report", self.raw)), 2)

    def test_missing_requirements_block(self):
        (self.lab / "requirements.yaml").unlink()
        self.assertTrue(quality.check_requirements(self.lab, "report", self.raw))

    def test_status_inside_table_is_found(self):
        table = self.document.add_table(rows=1, cols=1)
        table.cell(0,0).text = "[미제공: measurement]"
        self.document.save(self.docx)
        self.assertTrue(quality.check_docx(self.docx)["errors"])

    def test_wide_or_tall_image_is_found(self):
        path = self.ws / "tall.png"
        Image.new("RGB", (80, 500), "white").save(path)
        self.document.add_picture(str(path), width=Mm(200))
        self.document.save(self.docx)
        self.assertGreaterEqual(len(quality.check_docx(self.docx)["errors"]), 2)

    def test_image_fit_preserves_ratio_and_drawing_dimensions(self):
        from docx_post import fit_images
        path = self.ws / "tall.png"
        Image.new("RGB", (100, 400), "white").save(path)
        self.document.add_picture(str(path), width=Mm(100))
        self.assertEqual(fit_images(self.document, int(Mm(100)), int(Mm(180))), 1)
        extent = next(self.document.element.body.iter(qn("wp:extent")))
        self.assertAlmostEqual(int(extent.get("cy"))/int(extent.get("cx")), 4, places=4)
        for inner in self.document.element.body.iter(qn("a:ext")):
            self.assertEqual(inner.get("cx"), extent.get("cx"))
            self.assertEqual(inner.get("cy"), extent.get("cy"))

    def test_every_page_must_be_reviewed_once(self):
        self.create_render_fixture()
        for pages in ([1], [1,1,2], [1,2,3]):
            with self.assertRaises(ValueError):
                quality.review(self.docx, pages)
        quality.review(self.docx, [1,2])

    def test_changed_docx_invalidates_review(self):
        self.create_render_fixture()
        self.document.add_paragraph("changed")
        self.document.save(self.docx)
        with self.assertRaises(ValueError):
            quality.review(self.docx, [1,2])

    def test_changed_or_extra_page_invalidates_review(self):
        self.create_render_fixture()
        folder = quality.qa_dir(self.docx)
        Image.new("RGB", (90, 100), "black").save(folder / "page-1.png")
        with self.assertRaises(ValueError):
            quality.verify_render(self.docx)
        self.create_render_fixture()
        Image.new("RGB", (80, 100)).save(folder / "page-3.png")
        with self.assertRaises(ValueError):
            quality.verify_render(self.docx)

    def test_changed_pdf_invalidates_review(self):
        self.create_render_fixture()
        (quality.qa_dir(self.docx) / (self.docx.stem + ".pdf")).write_bytes(b"changed pdf")
        with self.assertRaises(ValueError):
            quality.verify_render(self.docx)

    def test_changed_source_invalidates_build(self):
        self.record_build()
        (self.lab / "source.txt").write_text("changed")
        with self.assertRaises(ValueError):
            quality.verify_build(self.docx, self.lab, "report")

    def test_old_engine_cannot_deliver_even_after_page_review(self):
        self.record_build()
        self.create_render_fixture()
        quality.review(self.docx, [1, 2])
        with patch.object(quality, 'engine_fingerprint', return_value={'scripts/math_layout.py': 'changed'}):
            with self.assertRaisesRegex(ValueError, '생성 엔진'):
                quality.deliver(self.docx, self.lab, 'report')
        record_path = quality.qa_dir(self.docx) / 'build.json'
        record = json.loads(record_path.read_text(encoding='utf-8'))
        del record['engine_sha256']
        quality.save_json(record_path, record)
        with self.assertRaisesRegex(ValueError, '버전 기록'):
            quality.verify_build(self.docx, self.lab, 'report')

    def test_missing_image_cannot_be_delivered_as_final(self):
        (self.lab / "report/report.md").write_text(self.raw + '\n![measurement](missing.png)\n', encoding="utf-8")
        self.record_build()
        with self.assertRaises(ValueError):
            quality.verify_build(self.docx, self.lab, "report")

    def test_malformed_requirements_report_an_error(self):
        (self.lab / "requirements.yaml").write_text('scope: [unterminated', encoding="utf-8")
        self.assertTrue(quality.check_requirements(self.lab, "report", self.raw))

    def test_delivery_requires_latest_review_and_copies_exact_docx(self):
        self.record_build()
        self.create_render_fixture()
        with self.assertRaises(FileNotFoundError):
            quality.deliver(self.docx, self.lab, "report")
        self.assertFalse((self.ws / "out/circuit").exists())
        quality.review(self.docx, [1,2])
        target = quality.deliver(self.docx, self.lab, "report")
        self.assertEqual(quality.digest(target), quality.digest(self.docx))

    def test_failed_renderer_invalidates_previous_review(self):
        self.create_render_fixture()
        quality.review(self.docx, [1,2])
        with patch.object(quality.subprocess, "run", side_effect=__import__('subprocess').CalledProcessError(1, "renderer")):
            with self.assertRaises(Exception):
                quality.render(self.docx, renderer=Path("fixture-renderer.py"))
        self.assertFalse((quality.qa_dir(self.docx) / "review.json").exists())
        self.assertFalse((quality.qa_dir(self.docx) / "render.json").exists())

    @unittest.skipUnless(shutil.which("pandoc"), "pandoc unavailable")
    def test_final_build_blocks_missing_requirements_before_docx(self):
        (self.lab / "requirements.yaml").unlink()
        self.docx.unlink()
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            build.build(self.lab, "report", final=True)
        self.assertFalse(self.docx.exists())

    @unittest.skipUnless(shutil.which("pandoc"), "pandoc unavailable")
    def test_real_build_draft_and_final_gate(self):
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            output = build.build(self.lab, "report", final=True)
        self.assertTrue(output.is_file())
        self.assertTrue((self.ws / "out/drafts/circuit" / output.name).is_file())
        self.assertFalse((self.ws / "out/circuit" / output.name).exists())
        headings = [p.text for p in Document(output).paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(headings, ["1. RC gain"])
        quality.verify_build(output, self.lab, "report")

    @unittest.skipUnless(shutil.which("pandoc"), "pandoc unavailable")
    def test_hw_delivery_preserves_code_and_packs_reviewed_pdf(self):
        import zipfile
        import pymupdf
        hw = self.ws / "courses/em/hw01"
        (hw / "code").mkdir(parents=True)
        (hw / "figs").mkdir()
        (hw / "meta.yaml").write_text('hw: 1\ntitle: Synthetic homework\n', encoding="utf-8")
        (hw / "source.txt").write_text("Synthetic Homework Problem 1", encoding="utf-8")
        (hw / "hw.md").write_text('# Problem 1) Synthetic exercise\n\nThe example result is described here.\n\n![Code](figs/p1_code.png){#fig:code}\n\n![Result](figs/p1_result.png){#fig:result}\n', encoding="utf-8")
        for name in ("p1_code.png", "p1_result.png"):
            Image.new("RGB", (200,100), "white").save(hw / "figs" / name)
        code_bytes = b'clc; clear;\n%% Problem 1\n% Synthetic test fixture\nx = 1;\n'
        (hw / "code/HW1.m").write_bytes(code_bytes)
        requirements = {"scope":{"hw":{"status":"confirmed", "source":"source.txt", "locator":"line 1"}},
            "items":[{"id":"1", "kind":"hw", "status":"answered", "source":"source.txt", "locator":"line 1", "answer":"Problem 1)", "evidence":["code/HW1.m", "figs/p1_code.png", "figs/p1_result.png"]}]}
        (hw / "requirements.yaml").write_text(yaml.safe_dump(requirements), encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            output = build.build(hw, "hw", final=True)
        self.docx = output
        self.create_render_fixture(count=1)
        pdf = quality.qa_dir(output) / (output.stem + ".pdf")
        document = pymupdf.open()
        document.new_page().insert_text((60,60), "Synthetic renderer fixture")
        document.save(pdf)
        document.close()
        manifest = json.loads((quality.qa_dir(output) / "render.json").read_text(encoding="utf-8"))
        manifest["pdf_sha256"] = quality.digest(pdf)
        quality.save_json(quality.qa_dir(output) / "render.json", manifest)
        quality.review(output, [1])
        with contextlib.redirect_stdout(io.StringIO()):
            target = quality.deliver(output, hw, "hw")
        self.assertEqual((target.parent / "HW1.m").read_bytes(), code_bytes)
        archive = target.with_suffix('.zip')
        with zipfile.ZipFile(archive) as zipped:
            self.assertEqual(zipped.read('HW1.m'), code_bytes)
            self.assertEqual(zipped.read(output.stem + '.pdf'), pdf.read_bytes())

    def test_duplicate_figure_ids_and_missing_references_are_warnings(self):
        style = yaml.safe_load((ROOT / "snu-report-core/templates/style.yaml").read_text(encoding="utf-8"))
        _, warnings = build.preprocess('![A](a.png){#fig:x}\n![B](b.png){#fig:x}\n@fig:absent', style)
        self.assertTrue(any("중복" in w for w in warnings))
        self.assertTrue(any("정의되지" in w for w in warnings))

    def test_page_break_insertion_is_idempotent(self):
        from docx_post import break_before_frames
        style = yaml.safe_load((ROOT / "snu-report-core/templates/style.yaml").read_text(encoding="utf-8"))
        table = self.document.add_table(rows=2, cols=1)
        table.cell(1,0).text = "Fig.1 - synthetic"
        self.document.save(self.docx)
        self.assertEqual(break_before_frames(self.docx, [("Fig",1)], style), 1)
        self.assertEqual(break_before_frames(self.docx, [("Fig",1)], style), 0)


if __name__ == "__main__":
    unittest.main()
