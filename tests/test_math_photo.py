"""Synthetic formula, figure and measurement-photo regressions."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import numpy as np
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Mm, Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'snu-report-core/scripts'))
sys.path.insert(0, str(ROOT/'snu-lab-photo/scripts'))
from math_layout import format_math, layout_issues, width_pt
from docx_post import clean_paragraphs, space_after_figures
from enhance_scope import enhance, rectify
from ooxml_order import normalize_document
import quality


def equation(p, text='x=1', fraction=False):
    xml = ('<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>'
           '<m:den><m:r><m:t>1+jωRC</m:t></m:r></m:den></m:f>' if fraction else
           '<m:r><m:t>'+text+'</m:t></m:r>')
    p._p.append(parse_xml('<m:oMathPara '+nsdecls('m','w')+'><m:oMath>'+xml+'</m:oMath></m:oMathPara>'))


class MathAndGapTests(unittest.TestCase):
    def test_fraction_growth_and_centering_preserve_equation_content(self):
        doc=Document()
        p=doc.add_paragraph()
        p.paragraph_format.line_spacing=Pt(6)
        p.paragraph_format.left_indent=Mm(25)
        equation(p, fraction=True)
        content=[t.text for t in p._p.iter(qn('m:t'))]
        format_math(doc,{})
        normalize_document(doc)
        self.assertEqual(content,[t.text for t in p._p.iter(qn('m:t'))])
        self.assertEqual(p.paragraph_format.left_indent,0)
        self.assertEqual(p._p.find('.//'+qn('m:jc')).get(qn('m:val')),'center')
        self.assertEqual(p._p.find('.//'+qn('w:spacing')).get(qn('w:lineRule')),'auto')
        self.assertTrue(p.paragraph_format.keep_together)
        self.assertFalse(layout_issues(doc))

    def test_inline_math_keeps_prose_alignment_but_removes_fixed_line_height(self):
        doc=Document()
        p=doc.add_paragraph('Measured voltage is ')
        p.paragraph_format.line_spacing=Pt(8)
        p._p.append(parse_xml('<m:oMath '+nsdecls('m')+'><m:r><m:t>V=2</m:t></m:r></m:oMath>'))
        format_math(doc,{})
        self.assertIsNone(p.alignment)
        self.assertEqual(p.paragraph_format.line_spacing,1.35)
        self.assertFalse(layout_issues(doc))

    def test_formula_table_row_expands_and_narrow_cell_is_flagged(self):
        doc=Document()
        table=doc.add_table(rows=1,cols=1)
        table.rows[0].height=Pt(6)
        table.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
        table.cell(0,0).width=Mm(18)
        equation(table.cell(0,0).paragraphs[0], '12345678901234567890')
        format_math(doc,{})
        self.assertEqual(table.rows[0].height_rule, WD_ROW_HEIGHT_RULE.AT_LEAST)
        self.assertTrue(any('폭' in x for x in layout_issues(doc)))

    def test_inline_math_overrides_inherited_exact_line_height(self):
        doc=Document()
        doc.styles['Normal'].paragraph_format.line_spacing=Pt(6)
        p=doc.add_paragraph('Value ')
        p.paragraph_format.space_after=Pt(4)
        p._p.append(parse_xml('<m:oMath '+nsdecls('m')+'><m:r><m:t>V</m:t></m:r></m:oMath>'))
        format_math(doc,{})
        self.assertEqual(p.paragraph_format.line_spacing,1.35)
        self.assertEqual(p.paragraph_format.space_after.pt,4)
        self.assertFalse(layout_issues(doc))

    def test_long_formula_is_rejected_by_delivery_check_until_split(self):
        doc=Document()
        equation(doc.add_paragraph(), 'x='+'1234567890+'*12)
        format_math(doc,{})
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'long.docx'
            doc.save(path)
            self.assertTrue(any('수식 폭' in x for x in quality.check_docx(path)['errors']))

    def test_fraction_width_does_not_add_stacked_numerator_and_denominator(self):
        doc=Document()
        p=doc.add_paragraph()
        equation(p,fraction=True)
        math=p._p.find('.//'+qn('m:oMath'))
        self.assertLess(width_pt(math), 70)

    def test_after_figure_gap_uses_paragraph_spacing_without_blank_enters(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'synthetic.png'
            Image.new('RGB',(40,30),'black').save(path)
            doc=Document()
            table=doc.add_table(rows=2,cols=1)
            table.cell(0,0).paragraphs[0].add_run().add_picture(str(path),width=Mm(40))
            table.cell(1,0).text='Fig.1 - synthetic'
            text=doc.add_paragraph('Interpretation after the figure.')
            clean_paragraphs(doc,{})
            space_after_figures(doc,{})
            self.assertEqual(len(doc.paragraphs),1)
            self.assertEqual(text.paragraph_format.space_before.pt,12)
            self.assertIsNone(table.cell(1,0).paragraphs[0].paragraph_format.space_before)


class PhotoTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.folder=Path(self.temp.name)
        self.source=self.folder/'source.png'
        # Uneven, gray illumination with bright red/green synthetic traces.
        a=np.tile(np.linspace(30,100,240,dtype=np.uint8),(120,1))
        rgb=np.stack([a,a,a],axis=-1)
        rgb[30:33,:,0]=220
        rgb[60:63,:,1]=220
        Image.fromarray(rgb).save(self.source)

    def tearDown(self):
        self.temp.cleanup()

    def test_identity_settings_preserve_pixels_source_and_provenance(self):
        before=hashlib.sha256(self.source.read_bytes()).hexdigest()
        target=self.folder/'identity.png'
        record=enhance(self.source,target,glare=0,contrast=1,sharpen=0)
        self.assertTrue(np.array_equal(np.asarray(Image.open(self.source)),np.asarray(Image.open(target))))
        self.assertEqual(record['source_sha256'],before)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(),before)
        self.assertEqual(json.loads(target.with_suffix('.png.json').read_text())['review_status'],'unreviewed')

    def test_overwrite_and_outside_crop_are_rejected(self):
        with self.assertRaises(ValueError):
            enhance(self.source,self.source)
        with self.assertRaises(ValueError):
            enhance(self.source,self.folder/'outside.png',crop=[-1,0,100,100])

    def test_smooth_glare_is_reduced_without_channel_swap_or_resize(self):
        target=self.folder/'readable.png'
        enhance(self.source,target,glare=.2,contrast=1,sharpen=0)
        a=np.asarray(Image.open(self.source))
        b=np.asarray(Image.open(target))
        self.assertEqual(a.shape,b.shape)
        self.assertLess(b[90,:,0].std(),a[90,:,0].std())
        self.assertTrue(np.all(b[31,:,0] > b[31,:,1]))
        self.assertTrue(np.all(b[61,:,1] > b[61,:,0]))

    def test_perspective_requires_observed_valid_quad(self):
        image=Image.open(self.source)
        result=rectify(image,[(10,10),(225,5),(230,110),(5,115)])
        self.assertLess(result.width,image.width)
        self.assertLess(result.height,image.height)
        with self.assertRaises(ValueError):
            rectify(image,[(0,0),(230,110),(230,0),(0,110)])


if __name__ == '__main__':
    unittest.main()
