#!/usr/bin/env python3
"""근거표(report/evidence.yaml) 검사·계산·표 생성

  python evidence.py <lab_dir>                      # 검사 + 요약 출력
  python evidence.py <lab_dir> --md report/evidence_tables.md   # 보고서에 붙일 표
  python evidence.py dump raw/data.xlsx             # 스프레드시트 셀 좌표·수식·값 덤프 (근거표 만들 때)

evidence.yaml 구조 (references/evidence.md 참고):
  methods:   측정 방법 (M01 …)       — 지연 측정 기준 등
  records:   원자료 (D01 …)           — kind: 측정 | 판독
  derived:   계산값 (C01 …)           — formula 로 다시 계산, compare_to 가 있으면 대조
  photos:    사진 (P01 …)
  questions: 사용자에게 확인할 것      — 보고서 결론을 바꾸는 것만

formula 에서 쓸 수 있는 것: mean(Dxx) min(Dxx) max(Dxx) n(Dxx) vals(Dxx)[i], + - * / ( ), 숫자, 다른 Cxx
"""
from __future__ import annotations

import _console  # noqa: F401  (Windows에서 한글·기호 출력)
import argparse
import math
import re
import sys
from pathlib import Path

import yaml

STATUS_RE = re.compile(r"\[(확인 필요|미제공|판독 불가|미측정|미수행|TODO)[^\]]*\]")


def decimals(x) -> int:
    s = f"{x}"
    return len(s.split(".")[1]) if "." in s and not s.endswith(".0") else 0


def fmt(x, d):
    return "?" if x is None else f"{x:.{d}f}"


def numeric(vals):
    return [float(v) for v in vals if isinstance(v, (int, float))]


def stats(rec):
    vals = rec.get("values", [])
    nums = numeric(vals)
    d = rec.get("decimals", max([1] + [decimals(v) for v in vals if isinstance(v, (int, float))]))
    if not nums:
        return {"n": 0, "mean": None, "min": None, "max": None, "d": d}
    return {"n": len(nums), "mean": sum(nums) / len(nums), "min": min(nums), "max": max(nums), "d": d}


def evaluate(formula, env_stats, env_derived):
    def get(key):
        if key not in env_stats:
            raise KeyError(f"없는 ID: {key}")
        return env_stats[key]

    ns = {
        "mean": lambda k: get(k)["mean"],
        "min": lambda k: get(k)["min"],
        "max": lambda k: get(k)["max"],
        "n": lambda k: get(k)["n"],
        "vals": lambda k: numeric(get(k)["raw"]),
        "sqrt": math.sqrt,
        "r1": lambda v: round(v + 1e-9, 1),  # 보고서에 적힌 소수 첫째 자리 값으로 계산
        "abs": abs,
    }
    expr = re.sub(r"\b(D\d+)\b", r'"\1"', formula)
    expr = re.sub(r"\b(C\d+)\b", lambda m: repr(env_derived.get(m.group(1))), expr)
    return eval(expr, {"__builtins__": {}}, ns)  # noqa: S307 (신뢰하는 로컬 파일)


def analyze(ev: dict):
    problems = []
    ids = set()
    rec_stats = {}
    for sec in ("methods", "records", "derived", "photos"):
        for r in ev.get(sec) or []:
            if r["id"] in ids:
                problems.append(f"중복 ID {r['id']}")
            ids.add(r["id"])
    for r in ev.get("records") or []:
        s = stats(r)
        s["raw"] = r.get("values", [])
        rec_stats[r["id"]] = s
        if r.get("kind") not in ("측정", "판독"):
            problems.append(f"{r['id']}: records의 kind는 측정/판독만 (지금: {r.get('kind')})")
        if "?" in [str(v) for v in r.get("values", [])]:
            problems.append(f"{r['id']}: 값 없음(?) 포함 → [미제공] 처리")
        if r.get("method") and r["method"] not in ids:
            problems.append(f"{r['id']}: 없는 방법 ID {r['method']}")
        for k, v in r.items():
            if isinstance(v, str) and STATUS_RE.search(v):
                problems.append(f"{r['id']}.{k}: {STATUS_RE.search(v).group(0)}")

    derived_vals = {}
    for c in ev.get("derived") or []:
        try:
            val = evaluate(c["formula"], rec_stats, derived_vals)
        except Exception as e:  # noqa: BLE001
            problems.append(f"{c['id']}: 계산 실패 ({e})")
            val = None
        derived_vals[c["id"]] = val
        cmp = c.get("compare_to")
        if cmp and val is not None:
            ref = float(cmp["value"])
            # 기본 허용오차 = 비교값이 적힌 자릿수의 반올림 오차 (25.28 → ±0.005)
            tol = float(cmp.get("tol", 0.5 * 10 ** -decimals(cmp["value"]) + 1e-9))
            if abs(val - ref) > tol:
                problems.append(
                    f"{c['id']}: 다시 계산한 값 {val:.3f} ≠ {cmp.get('source', '비교값')} {ref} (차이 {val - ref:+.3f}) → [확인 필요]")
        for k, v in c.items():
            if isinstance(v, str) and STATUS_RE.search(v):
                problems.append(f"{c['id']}.{k}: {STATUS_RE.search(v).group(0)}")
    for q in ev.get("questions") or []:
        problems.append(f"질문: {q}")
    return rec_stats, derived_vals, problems


def to_md(ev, rec_stats, derived_vals) -> str:
    L = []
    recs = ev.get("records") or []
    if recs:
        maxn = max(len(r.get("values", [])) for r in recs)
        head = ["ID", "항목", "조건", "구분"] + [f"#{i + 1}" for i in range(maxn)] + ["평균", "단위"]
        L += ["Table: 측정 원자료와 평균 {#tbl:raw}", "",
              "| " + " | ".join(head) + " |", "|" + "|".join([":-:"] * len(head)) + "|"]
        for r in recs:
            s = rec_stats[r["id"]]
            vals = [str(v) for v in r.get("values", [])] + [""] * (maxn - len(r.get("values", [])))
            L.append("| " + " | ".join([r["id"], r.get("item", ""), r.get("condition", ""), r.get("label", "")]
                                         + vals + [fmt(s["mean"], s["d"]), r.get("unit", "")]) + " |")
        L.append("")
    der = ev.get("derived") or []
    if der:
        L += ["Table: 계산값 {#tbl:derived}", "",
              "| ID | 내용 | 식 | 값 | 단위 |", "|:-:|:--|:--|:-:|:-:|"]
        for c in der:
            v = derived_vals.get(c["id"])
            d = c.get("decimals", 2)
            L.append(f"| {c['id']} | {c.get('what', '')} | `{c['formula']}` | {fmt(v, d)} | {c.get('unit', '')} |")
        L.append("")
    return "\n".join(L)


def dump(path: Path):
    import openpyxl

    wb_f = openpyxl.load_workbook(path)                 # 수식
    wb_v = openpyxl.load_workbook(path, data_only=True)  # 캐시된 값 (없으면 None)
    for ws in wb_f:
        wv = wb_v[ws.title]
        print(f"### {path.name} :: {ws.title} ({ws.dimensions})")
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None:
                    continue
                cached = wv[c.coordinate].value
                if isinstance(c.value, str) and c.value.startswith("="):
                    cells.append(f"{c.coordinate}={c.value} → {cached}")
                else:
                    cells.append(f"{c.coordinate}:{c.value!r}")
            if cells:
                print("  " + " | ".join(cells))


def main():
    if len(sys.argv) >= 3 and sys.argv[1] == "dump":
        for p in sys.argv[2:]:
            dump(Path(p))
        return
    ap = argparse.ArgumentParser()
    ap.add_argument("lab_dir", type=Path)
    ap.add_argument("--md")
    a = ap.parse_args()
    p = a.lab_dir / "report" / "evidence.yaml"
    ev = yaml.safe_load(p.read_text(encoding="utf-8"))
    rec_stats, derived_vals, problems = analyze(ev)

    print("## 원자료")
    for r in ev.get("records") or []:
        s = rec_stats[r["id"]]
        rng = f"{s['min']}–{s['max']}" if s["n"] else "-"
        print(f"  {r['id']} {r.get('item', '')} {r.get('label', '')} {r.get('condition', '')}: "
              f"n={s['n']} 평균 {fmt(s['mean'], s['d'])} {r.get('unit', '')} (범위 {rng})")
    print("## 계산값")
    for c in ev.get("derived") or []:
        print(f"  {c['id']} {c.get('what', '')}: {fmt(derived_vals.get(c['id']), c.get('decimals', 2))} {c.get('unit', '')}")
    if problems:
        print("## 확인할 것")
        for x in problems:
            print("  ⚠", x)
    if a.md:
        Path(a.md).write_text(to_md(ev, rec_stats, derived_vals), encoding="utf-8")
        print("✓", a.md)


if __name__ == "__main__":
    main()
