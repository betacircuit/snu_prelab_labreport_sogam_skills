#!/usr/bin/env python3
"""보고서 구조, 문항 대응, 실제 렌더링 파일과 전 쪽 검토를 연결한다.

검사는 내용의 진실성이나 시각적 품질을 자동으로 보증하지 않는다.
render -> 모든 PNG를 실제로 읽기 -> review --pages 1,2,... -> deliver 순서다.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import _console  # noqa: F401
from docx import Document
from docx.oxml.ns import qn
from ws import find_workspace

STATUS = re.compile(r"\[(?:TODO|미제공|확인 필요|판독 불가|미측정|미수행)(?::[^\]]*)?\]")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def engine_fingerprint() -> dict:
    """A render/review from an older generator must not certify a current delivery."""
    root = Path(__file__).resolve().parent.parent
    paths = sorted((root / 'scripts').glob('*.py')) + sorted((root / 'scripts').glob('*.ps1'))
    paths += sorted((root / 'templates').glob('*.yaml'))
    return {p.relative_to(root).as_posix(): digest(p) for p in paths}


def qa_dir(docx: Path) -> Path:
    return docx.parent / "qa" / docx.stem


def save_json(path: Path, data: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def record_build(docx: Path, lab_dir: Path, kind: str, style_path: Path):
    """원고·설정·문항 원문·그림을 생성 시점에 묶어 둔다."""
    import yaml
    from build import SRC
    from ws import course_skills
    raw_path = lab_dir / SRC[kind]
    paths = [raw_path, lab_dir / "meta.yaml", style_path]
    ws = find_workspace(lab_dir)
    if ws:
        paths += [ws / "profile.yaml", lab_dir.parent / "course.yaml"]
    default_course = course_skills().get(lab_dir.parent.name)
    if default_course:
        paths.append(default_course / "course.yaml")
    if kind == "hw":
        from build import load_vars
        paths.append(lab_dir / "code" / f"HW{load_vars(lab_dir, kind)['lab_num']}.m")
    requirement = lab_dir / "requirements.yaml"
    paths.append(requirement)
    if requirement.exists():
        data = yaml.safe_load(requirement.read_text(encoding="utf-8")) or {}
        if isinstance(data, dict):
            scopes = data.get("scope", {})
            records = list(scopes.values()) if isinstance(scopes, dict) else []
            records += data.get("items", []) if isinstance(data.get("items", []), list) else []
            for entry in records:
                if isinstance(entry, dict):
                    sources = [entry.get("source")] + (entry.get("evidence", []) if isinstance(entry.get("evidence", []), list) else [])
                    paths += [lab_dir / s for s in sources if isinstance(s, str) and not s.startswith("conversation:")]
    missing_images = []
    for image in re.findall(r"!\[[^\]]*\]\(([^)]+)\)", raw_path.read_text(encoding="utf-8")):
        p = Path(image)
        resolved = p if p.is_absolute() else next((d/p for d in (raw_path.parent, lab_dir) if (d/p).exists()), raw_path.parent/p)
        paths.append(resolved)
        if not resolved.is_file():
            missing_images.append(image)
    inputs = {str(p.resolve()): digest(p) if p.is_file() else None for p in paths}
    save_json(qa_dir(docx) / "build.json", {"kind": kind, "lab_dir": str(lab_dir.resolve()),
              "docx_sha256": digest(docx), "inputs": inputs, "missing_images": missing_images,
              "engine_sha256": engine_fingerprint()})


def verify_build(docx: Path, lab_dir: Path, kind: str):
    record = json.loads((qa_dir(docx) / "build.json").read_text(encoding="utf-8"))
    if record.get('engine_sha256') != engine_fingerprint():
        raise ValueError('생성 엔진이 변경됐거나 버전 기록이 없음 — 최신 엔진으로 다시 빌드·렌더링·검토')
    if record.get("missing_images"):
        raise ValueError("누락 사진이 있는 초안 — 자료를 넣고 다시 빌드")
    if record.get("kind") != kind or record.get("lab_dir") != str(lab_dir.resolve()) or record.get("docx_sha256") != digest(docx):
        raise ValueError("빌드 기록과 DOCX/과제가 다름 — 원고에서 다시 빌드")
    for name, value in record["inputs"].items():
        p = Path(name)
        if (digest(p) if p.is_file() else None) != value:
            raise ValueError("빌드 뒤 원고·설정·근거가 변경됨: " + name)


def check_docx(path: Path) -> dict:
    errors, warnings = [], []
    with ZipFile(path) as z:
        if z.testzip():
            errors.append("DOCX zip 손상")
    doc = Document(path)
    from math_layout import layout_issues
    errors.extend(layout_issues(doc))
    body = doc.element.body
    text = "\n".join(t.text or "" for t in body.iter(qn("w:t")))
    if STATUS.search(text) or "사진 넣을 곳" in text or '⟦미제공:' in text:
        errors.append("초안 상태 표시 또는 누락 사진 칸이 남음")
    if re.search(r"\?\?(?:fig|tbl):|@(?:fig|tbl):|\{\{[^}]+\}\}", text):
        errors.append("미해결 그림/표 참조 또는 템플릿 변수가 남음")
    from proof import deduplicate_heading_labels
    for paragraph in doc.paragraphs:
        if (paragraph.style.style_id.startswith('Heading')
                and deduplicate_heading_labels(paragraph.text) != paragraph.text):
            errors.append('제목 번호 중복: ' + paragraph.text)
    for element in body:
        if element.find('.//' + qn('w:drawing')) is None:
            continue
        following = element.getnext()
        if following is not None:
            ps = following.find('./' + qn('w:pPr') + '/' + qn('w:pStyle'))
            if ps is not None and ps.get(qn('w:val')) == 'ImageCaption':
                following = following.getnext()
        if following is None or following.tag == qn('w:sectPr'):
            continue
        if (following.tag != qn('w:p') or ''.join(following.itertext()).strip()
                or following.find('.//' + qn('w:drawing')) is not None):
            errors.append('그림·캡션 뒤 실제 빈 문단(Enter)이 없음')
    sizes = [(s.page_width-s.left_margin-s.right_margin,
              s.page_height-s.top_margin-s.bottom_margin) for s in doc.sections]
    max_width, max_height = max(x[0] for x in sizes), max(x[1] for x in sizes)
    for extent in body.iter(qn("wp:extent")):
        if int(extent.get("cx", 0)) > max_width + 12700:
            errors.append("그림이 본문 폭을 넘음")
        if int(extent.get("cy", 0)) > max_height - 350000:
            errors.append("그림과 캡션을 한 쪽에 배치할 높이가 부족함")
    for table in doc.tables:
        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            width = sum(int(c.get(qn("w:w"), 0)) for c in grid)
            if width * 635 > max_width + 12700:
                errors.append("표가 본문 폭을 넘음")
    return {"docx": str(path.resolve()), "sha256": digest(path),
            "errors": sorted(set(errors)), "warnings": sorted(set(warnings))}


def check_requirements(lab_dir: Path, kind: str, manuscript: str) -> list[str]:
    import yaml
    path = lab_dir / "requirements.yaml"
    if not path.exists():
        return ["requirements.yaml이 없음 — 원문 문항 대응 미검증"]
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except yaml.YAMLError as error:
        return ["requirements.yaml 문법 오류: " + str(error)]
    if not isinstance(data, dict):
        return ["requirements.yaml은 mapping이어야 함"]
    errors = []

    def source_valid(record, label):
        source, locator = record.get("source"), record.get("locator")
        if not isinstance(source, str) or not source.strip() or not str(locator or "").strip():
            errors.append(f"{label}: 원문 파일/대화와 위치가 없음")
        elif not source.startswith("conversation:") and not (lab_dir / source).is_file():
            errors.append(f"{label}: 원문 파일 없음: {source}")

    scopes = data.get("scope", {})
    scope = scopes.get(kind, {}) if isinstance(scopes, dict) else {}
    if not isinstance(scope, dict):
        scope = {}
    if scope.get("status") != "confirmed":
        errors.append(f"{kind}: 원문 범위가 confirmed가 아님")
    source_valid(scope, kind + " 범위")
    raw_items = data.get("items", [])
    if not isinstance(raw_items, list) or any(not isinstance(i, dict) for i in raw_items):
        return errors + ["items는 문항 mapping 목록이어야 함"]
    items = [i for i in raw_items if i.get("kind") == kind]
    if not items:
        errors.append(f"{kind}: 대응 문항이 없음")
    seen = set()
    headings = [re.sub(r"\s*\{[^}]*\}\s*$", "", line.lstrip("# ")).strip()
                for line in manuscript.splitlines() if re.match(r"^#{1,6}\s", line)]
    for item in items:
        label = str(item.get("id") or "(번호 없음)")
        if not item.get("id"):
            errors.append("문항 번호 id가 없음")
        if label in seen:
            errors.append(f"중복 문항: {label}")
        seen.add(label)
        source_valid(item, label)
        answer = item.get("answer")
        if not isinstance(answer, str) or not answer.strip() or not any(answer.strip() in h for h in headings):
            errors.append(f"{label}: 실제 원고의 답변 제목을 찾을 수 없음")
        if item.get("status") != "answered":
            errors.append(f"{label}: 답변 상태가 answered가 아님")
        evidence = item.get("evidence", [])
        if not isinstance(evidence, list):
            errors.append(f"{label}: evidence는 파일 경로 목록이어야 함")
            continue
        for source in evidence:
            if not isinstance(source, str) or not (lab_dir / source).is_file():
                errors.append(f"{label}: 근거 파일 없음: {source}")
    return errors


def render(docx: Path, renderer: Path | None = None, soffice: str | None = None, word: bool = False) -> dict:
    if word and (renderer or soffice):
        raise ValueError('--word는 --renderer/--soffice와 함께 사용할 수 없음')
    if word and os.name != 'nt':
        raise ValueError('--word 렌더링은 Microsoft Word가 설치된 Windows에서만 가능')
    folder = qa_dir(docx)
    folder.mkdir(parents=True, exist_ok=True)
    # 이전 렌더링은 새 파일로 오인하지 않도록 무효화하고 이 폴더의 쪽 PNG만 정리.
    for name in ("render.json", "review.json"):
        (folder / name).unlink(missing_ok=True)
    for page in folder.glob("page-*.png"):
        page.unlink()
    (folder / (docx.stem + ".pdf")).unlink(missing_ok=True)
    before = digest(docx)
    env = os.environ.copy()
    if soffice:
        env["SNU_SOFFICE"] = soffice
        env["PATH"] = str(Path(soffice).parent) + os.pathsep + env.get("PATH", "")
    if word:
        from ws import find_tool
        import pymupdf
        powershell = find_tool('pwsh') or find_tool('powershell')
        if not powershell:
            raise ValueError('Word 렌더링용 PowerShell을 찾지 못함')
        pdf = folder / (docx.stem + '.pdf')
        subprocess.run([powershell, '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                        '-File', str(Path(__file__).with_name('render_word.ps1')),
                        '-InputDocx', str(docx.resolve()), '-OutputPdf', str(pdf.resolve())],
                       check=True, timeout=180)
        with pymupdf.open(pdf) as document:
            for i, page in enumerate(document):
                page.get_pixmap(dpi=150).save(folder / f'page-{i+1}.png')
    elif renderer:
        subprocess.run([sys.executable, str(renderer.resolve()), str(docx.resolve()),
                        "--output_dir", str(folder), "--emit_pdf"], check=True, env=env, timeout=180)
    else:
        from build import to_pdf
        import pymupdf
        pdf = to_pdf(docx, output_dir=folder, soffice=soffice)
        with pymupdf.open(pdf) as document:
            for i, page in enumerate(document):
                page.get_pixmap(dpi=150).save(folder / f"page-{i+1}.png")
    if before != digest(docx):
        raise ValueError("렌더링 도중 DOCX가 변경됨 — 다시 렌더링")
    pages = sorted(folder.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    numbers = [int(p.stem.split("-")[-1]) for p in pages]
    if numbers != list(range(1, len(pages)+1)) or not pages or any(p.stat().st_size == 0 for p in pages):
        raise ValueError("연속된 전체 쪽 PNG가 생성되지 않음")
    from PIL import Image
    for page in pages:
        with Image.open(page) as im:
            im.verify()
    record = {"docx_sha256": before, "pages": {str(i+1): digest(p) for i,p in enumerate(pages)},
              "renderer": 'Microsoft Word' if word else str(renderer or 'LibreOffice')}
    pdf = folder / (docx.stem + ".pdf")
    record["pdf_sha256"] = digest(pdf) if pdf.is_file() else None
    save_json(folder / "render.json", record)
    return record


def verify_render(docx: Path) -> dict:
    folder = qa_dir(docx)
    record = json.loads((folder / "render.json").read_text(encoding="utf-8"))
    if digest(docx) != record["docx_sha256"]:
        raise ValueError("DOCX가 렌더링 뒤 변경됨")
    if {p.name for p in folder.glob("page-*.png")} != {f"page-{n}.png" for n in record["pages"]}:
        raise ValueError("렌더링 쪽 목록이 변경됨")
    for n, value in record["pages"].items():
        if digest(folder / f"page-{n}.png") != value:
            raise ValueError("렌더링 이미지가 변경됨")
    pdf = folder / (docx.stem + ".pdf")
    if (digest(pdf) if pdf.is_file() else None) != record.get("pdf_sha256"):
        raise ValueError("렌더링 PDF가 변경됨")
    return record


def review(docx: Path, pages: list[int]) -> dict:
    record = verify_render(docx)
    expected = sorted(map(int, record["pages"]))
    if sorted(pages) != expected:
        raise ValueError("실제로 확인한 전체 쪽 번호를 중복 없이 명시해야 함")
    record["reviewed_pages"] = expected
    save_json(qa_dir(docx) / "review.json", record)
    return record


def deliver(docx: Path, lab_dir: Path, kind: str) -> Path:
    from build import SRC, load_vars, submission_stem
    verify_build(docx, lab_dir, kind)
    result = check_docx(docx)
    raw = (lab_dir / SRC[kind]).read_text(encoding="utf-8")
    result["errors"] += check_requirements(lab_dir, kind, raw)
    if result["errors"]:
        raise ValueError("최종 전달 불가:\n  " + "\n  ".join(result["errors"]))
    rendered = verify_render(docx)
    reviewed = json.loads((qa_dir(docx) / "review.json").read_text(encoding="utf-8"))
    if reviewed.get("docx_sha256") != rendered["docx_sha256"] or reviewed.get("pages") != rendered["pages"] or reviewed.get("pdf_sha256") != rendered.get("pdf_sha256") or reviewed.get("reviewed_pages") != sorted(map(int, rendered["pages"])):
        raise ValueError("최신 DOCX의 전 쪽 검토가 완료되지 않음")
    if kind == "hw" and rendered.get("pdf_sha256") is None:
        raise ValueError("기전연 제출 zip에 필요한 검토 PDF가 없음")
    ws = find_workspace(lab_dir)
    if ws is None:
        raise ValueError("profile.yaml 작업 폴더를 찾지 못함")
    values = load_vars(lab_dir, kind)
    if kind == "hw" and not (lab_dir / "code" / f"HW{values['lab_num']}.m").is_file():
        raise ValueError("기전연 제출에 필요한 본인 MATLAB 코드가 없음")
    expected = submission_stem(kind, values) + ".docx"
    if docx.name != expected:
        raise ValueError(f"제출 파일명 불일치: {expected}")
    target = ws / "out" / str(values.get("key") or lab_dir.parent.name) / expected
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(docx, target)
    if digest(target) != rendered["docx_sha256"]:
        raise ValueError("전달 사본이 검토한 DOCX와 다름")
    save_json(qa_dir(docx) / "delivery.json", {**result, "delivered": str(target)})
    if kind == "hw":
        # 검토한 DOCX에서 생성된 PDF만 제출 zip으로 묶는다.
        pdf = qa_dir(docx) / (docx.stem + ".pdf")
        if pdf.is_file():
            from mcode import pack
            pack(lab_dir, pdf)
            shutil.copy2(pdf, target.parent / pdf.name)
            code = lab_dir / "code" / f"HW{values['lab_num']}.m"
            shutil.copy2(code, target.parent / code.name)
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["check", "render", "review", "deliver"])
    parser.add_argument("docx", type=Path)
    parser.add_argument("--renderer", type=Path)
    parser.add_argument("--soffice")
    parser.add_argument('--word', action='store_true', help='Windows Microsoft Word로 읽기 전용 PDF/PNG 렌더링')
    parser.add_argument("--pages", help="실제로 이미지 확인한 모든 쪽 번호: 1,2,3")
    parser.add_argument("--lab-dir", type=Path)
    parser.add_argument("--kind", choices=["prelab", "report", "hw"])
    args = parser.parse_args()
    try:
        if args.action == "check":
            result = check_docx(args.docx)
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return bool(result["errors"])
        if args.action == "render":
            render(args.docx, args.renderer, args.soffice, args.word)
            print("렌더링 완료, 전 쪽 이미지 검토 필요:", qa_dir(args.docx))
        elif args.action == "review":
            if not args.pages:
                parser.error("--pages가 필요함")
            review(args.docx, [int(n) for n in args.pages.split(",")])
            print("전 쪽 검토 기록 완료")
        else:
            if args.lab_dir is None or args.kind is None:
                parser.error("--lab-dir와 --kind가 필요함")
            print("최종 전달:", deliver(args.docx, args.lab_dir.resolve(), args.kind))
    except (OSError, ValueError, KeyError, RuntimeError, ImportError, subprocess.SubprocessError) as error:
        print("검증 실패:", error, file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
