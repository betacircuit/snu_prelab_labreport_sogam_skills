#!/usr/bin/env python3
"""수업 자료 자동 분류: inbox/ 에 던져 둔 파일을 [Lab NN] 번호로 정리한다.

  python ingest.py                                   # inbox/<과목>/* 전부 처리 (logic, circuit …)
  python ingest.py --course logic                    # 한 과목만
  python ingest.py --course logic 파일1 파일2 ...     # 특정 파일 처리 (원본은 복사, 삭제 안 함)
  python ingest.py --course logic --lab 01 측정값.xlsx  # 파일명에 Lab 번호가 없을 때 직접 지정
  python ingest.py --dry-run                         # 어디로 갈지 미리 보기

과목 폴더: courses/<과목>/ (course.yaml이 있는 폴더). inbox/<과목>/에 올린 파일을 그 과목으로 분류한다.

분류 규칙
  Lab 번호는 이 순서로 찾는다 (찾은 근거를 출력한다)
    1. 파일명: "[Lab 01] ...", "Lab01_...", "lab01", "실험 3", "실험3_", "제3실험", "Experiment 3", "Exp.3",
               "예비보고서 3", "결과보고서3"
    2. 파일 내용: PDF·docx·pptx·txt 앞부분의 "Lab 3", "실험 3", "Experiment 3" (제목 쪽을 우선, 없으면 가장 많이 나온 번호)
    3. 날짜: 파일명·내용·사진 촬영일의 날짜가 이미 있는 labNN/meta.yaml의 실험일(lab_date)과 맞으면 그 Lab
    4. 그래도 못 찾으면 inbox/<과목>/_unsorted/ — 에이전트가 파일을 직접 읽고 주제·날짜로 Lab을 정해
       `--lab NN`으로 다시 돌린다 (사용자에게 묻지 않는다. course.md 일정표의 주제와 대조)
  Lab 00                → courses/<과목>/materials/          (공통 규칙: 보고서 가이드라인, 오리엔테이션)
  Lab NN                → courses/<과목>/labNN/materials/    (guidebook.pdf, slides_YYMMDD.pdf 등)
  prelabNN_학번_이름.*   → courses/<과목>/labNN/submitted/    (내가 낸 파일 사본)
  labNN_학번_이름.*      → courses/<과목>/labNN/submitted/
  .xlsx/.csv 등 데이터   → courses/<과목>/labNN/report/raw/
  이미지(.jpg/.png/.heic) → courses/<과목>/labNN/report/photos/
  번호 인식 실패         → inbox/<과목>/_unsorted/ 에 남기고 알려 줌

같은 내용(sha256) 파일은 한 번만 저장. PDF는 옆에 .txt(텍스트 추출)를 만들어 Claude가 바로 읽을 수 있게 한다.
Lab NN 자료가 새로 들어오면 courses/<과목>/labNN/meta.yaml, requirements.md 초안을 만든다 (이미 있으면 건드리지 않음).
"""
from __future__ import annotations

import _console  # noqa: F401  (Windows에서 한글·기호 출력)
import argparse
import datetime as dt
import hashlib
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
from ws import course_defaults, find_tool, find_workspace, known_courses, remember_engine  # noqa: E402

PROFILE: dict = {}
UPLOAD_PREFIX = re.compile(r"^[0-9a-f]{8}-")  # claude 업로드가 붙이는 접두어
LAB_RE = re.compile(r"(?i)(?<![a-z])(pre)?\s*lab\s*[_\- .]?\s*(\d{1,2})(?!\d)")
# "실험 3", "실험3", "제3실험", "제 3 장 실험", "Experiment 3", "Exp. 3", "예비보고서 3", "결과 보고서3"
LAB_ALT_RES = [
    re.compile(r"(?<![가-힣])실험\s*[_\- #]?\s*(\d{1,2})(?![\d.])"),
    re.compile(r"제\s*(\d{1,2})\s*(?:실험|장)"),
    re.compile(r"(?i)(?<![a-z])exp(?:eriment)?\.?\s*[_\- #]?\s*(\d{1,2})(?!\d)"),
    re.compile(r"(?:예비|결과)\s*보고서\s*[_\- #]?\s*(\d{1,2})(?!\d)"),
]
SUBMIT_RE = re.compile(r"(?i)^(pre)?lab(\d{1,2})_(\d{4}-\d{5})_")
DATE6 = re.compile(r"(?<!\d)(2\d)(0[1-9]|1[0-2])([0-3]\d)(?!\d)")
DATA_EXT = {".xlsx", ".xls", ".csv", ".tsv", ".json"}
IMG_EXT = {".jpg", ".jpeg", ".png", ".heic", ".bmp", ".gif", ".webp"}
MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july", "august",
     "september", "october", "november", "december"], 1)}


# ───────────────────────── helpers ─────────────────────────
def find_root(start: Path) -> Path:
    return find_workspace(start) or start


def sha(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def clean_name(name: str) -> str:
    name = UPLOAD_PREFIX.sub("", name)
    name = re.sub(r"_(\d)(?=\.\w+$)", "", name)  # "파일_1.pdf" 같은 중복 다운로드 꼬리표
    return name


def lab_from_name(stem: str):
    """파일명 → (lab 'NN' | None, 근거)"""
    m = LAB_RE.search(stem)
    if m:
        return f"{int(m.group(2)):02d}", f"파일명 '{m.group(0).strip()}'"
    for r in LAB_ALT_RES:
        m = r.search(stem)
        if m:
            return f"{int(m.group(1)):02d}", f"파일명 '{m.group(0).strip()}'"
    return None, ""


def head_text(path: Path, limit: int = 6000) -> str:
    """파일 앞부분 글자 (PDF 첫 3쪽, docx 앞 문단, pptx 앞 슬라이드, 텍스트)"""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            try:
                import pymupdf
                with pymupdf.open(str(path)) as doc:
                    return "\n".join(doc[i].get_text() for i in range(min(3, len(doc))))[:limit]
            except ImportError:
                if find_tool("pdftotext"):
                    r = subprocess.run([find_tool("pdftotext"), "-l", "3", "-layout", str(path), "-"],
                                       capture_output=True, text=True, check=False)
                    return r.stdout[:limit]
        elif ext == ".docx":
            from docx import Document
            return "\n".join(p.text for p in Document(str(path)).paragraphs[:80])[:limit]
        elif ext == ".pptx":
            import zipfile
            with zipfile.ZipFile(path) as z:
                names = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                               key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))[:4]
                xml = " ".join(z.read(n).decode("utf-8", "ignore") for n in names)
                return " ".join(re.findall(r"<a:t>([^<]*)</a:t>", xml))[:limit]
        elif ext in (".txt", ".md", ".csv", ".tsv"):
            return path.read_text(encoding="utf-8", errors="ignore")[:limit]
        elif ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            cells = [wb.sheetnames and " ".join(wb.sheetnames)]
            for ws in wb.worksheets[:3]:
                for row in ws.iter_rows(max_row=15, values_only=True):
                    cells += [str(c) for c in row if c is not None]
            return " ".join(cells)[:limit]
    except Exception:
        return ""
    return ""


def lab_from_content(path: Path):
    """파일 앞부분 글자 → (lab | None, 근거). 제목 쪽(앞 600자)에 나온 번호를 우선하고,
    없으면 두 번 이상 나오면서 가장 많이 나온 번호."""
    text = head_text(path)
    if not text:
        return None, ""
    pats = [re.compile(r"(?i)(?<![a-z])(?:pre\s*)?lab(?:oratory)?\s*[_\- .#]?\s*(\d{1,2})(?!\d)")] + LAB_ALT_RES
    hits = []
    for r in pats:
        for m in r.finditer(text):
            n = int(m.group(m.lastindex))
            if 0 <= n <= 20:
                hits.append((m.start(), n, m.group(0).strip()))
    if not hits:
        return None, ""
    hits.sort()
    head = [h for h in hits if h[0] < 600]
    if head:
        return f"{head[0][1]:02d}", f"내용 앞부분 '{head[0][2]}'"
    from collections import Counter
    cnt = Counter(n for _, n, _ in hits).most_common(2)
    if cnt[0][1] >= 2 and (len(cnt) == 1 or cnt[0][1] > cnt[1][1]):
        return f"{cnt[0][0]:02d}", f"내용에 'Lab/실험 {cnt[0][0]}' {cnt[0][1]}번"
    return None, ""


DATE8 = re.compile(r"(?<!\d)(20\d{2})[-_.]?(0[1-9]|1[0-2])[-_.]?([0-3]\d)(?!\d)")
KO_DATE = re.compile(r"(?:(20\d{2})\s*년\s*)?(\d{1,2})\s*월\s*(\d{1,2})\s*일")


def dates_of(path: Path, text: str = "") -> list[dt.date]:
    out = []
    name = clean_name(path.name)
    for m in DATE8.finditer(name):
        out.append((int(m.group(1)), int(m.group(2)), int(m.group(3))))
    for m in DATE6.finditer(name):
        out.append((2000 + int(m.group(1)), int(m.group(2)), int(m.group(3))))
    year = dt.date.today().year
    for m in KO_DATE.finditer(text[:3000]):
        out.append((int(m.group(1) or year), int(m.group(2)), int(m.group(3))))
    d = _date_from_text(text[:3000], year) if text else None
    if d:
        out.append((d.year, d.month, d.day))
    if path.suffix.lower() in IMG_EXT:
        try:
            from PIL import Image
            ex = Image.open(path).getexif()
            raw = ex.get(36867) or ex.get(306) or (ex.get_ifd(0x8769).get(36867) if hasattr(ex, "get_ifd") else None)
            if raw:
                y, mo, da = re.match(r"(\d{4}):(\d{2}):(\d{2})", str(raw)).groups()
                out.append((int(y), int(mo), int(da)))
        except Exception:
            pass
    res = []
    for y, mo, da in out:
        try:
            res.append(dt.date(y, mo, da))
        except ValueError:
            pass
    return res


def lab_from_date(path: Path, croot: Path, text: str = ""):
    """날짜가 이미 있는 Lab의 실험일과 맞으면 그 Lab (사진·측정 파일은 실험 당일, 슬라이드는 그 주)"""
    labs = {}
    for meta in croot.glob("lab*/meta.yaml"):
        try:
            d = yaml.safe_load(meta.read_text(encoding="utf-8")) or {}
            labs[meta.parent.name[3:]] = dt.date.fromisoformat(str(d.get("lab_date"))[:10])
        except Exception:
            continue
    if not labs:
        return None, ""
    for d in dates_of(path, text):
        same = [k for k, v in labs.items() if v == d]
        if len(same) == 1:
            return same[0], f"날짜 {d} = lab{same[0]} 실험일"
        week = [k for k, v in labs.items() if 0 <= (v - d).days <= 6]
        if len(week) == 1:
            return week[0], f"날짜 {d} → lab{week[0]} 실험일 {labs[week[0]]} 주"
    return None, ""


def slug(s: str) -> str:
    s = re.sub(r"(?i)\[?\s*(pre)?lab\s*[_\- .]?\d{1,2}\s*\]?", "", s)
    s = re.sub(r"(?i)\[?\s*(?:실험|exp(?:eriment)?\.?|(?:예비|결과)\s*보고서)\s*[_\- #]?\s*\d{1,2}\s*\]?", "", s)
    s = re.sub(r"[^\w가-힣]+", "_", s).strip("_").lower()
    return s or "file"


HW_RES = [re.compile(r"(?i)(?<![a-z])(?:hw|homework|assignment)\s*[_\- #.]?\s*(\d{1,2})(?!\d)"),
          re.compile(r"(?<![가-힣])(?:과제|실습)\s*[_\- #]?\s*(\d{1,2})(?!\d)")]   # 기전연 자료는 '실습1' = HW1


def classify_hw(path: Path, forced: str | None):
    """과제(HW) 단위 과목: HW 번호 → (NN | None, 하위 폴더, 파일명, 근거). .m은 code/, 나머지는 materials/"""
    name = clean_name(path.name)
    stem, ext = Path(name).stem, Path(name).suffix.lower()
    num, why = (forced, "--hw") if forced else (None, "")
    if num is None:
        for r in HW_RES:
            m = r.search(stem)
            if m:
                num, why = f"{int(m.group(1)):02d}", f"파일명 '{m.group(0).strip()}'"
                break
    if num is None:
        text = head_text(path)
        for r in HW_RES:
            m = r.search(text[:800])
            if m:
                num, why = f"{int(m.group(1)):02d}", f"내용 앞부분 '{m.group(0).strip()}'"
                break
    if num is None:
        return None, None, name, ""
    if re.fullmatch(r"(?i)HW\d+_.+_\d{4}-\d{5}", stem):
        return num, "submitted", name, why
    if ext == ".m":
        return num, "code", f"HW{int(num)}.m" if re.fullmatch(r"(?i)hw\d+", stem) else name, why
    if ext in IMG_EXT:   # 사용자 MATLAB 캡처 (코드 화면, 결과 그림) — 컬러 그대로 보고서에 넣는다
        return num, "figs", name, why
    return num, "materials", name, why


def classify(path: Path, forced_lab: str | None, croot: Path | None = None):
    """→ (lab_num | None, 대상 하위 경로(폴더), 파일명, 근거)"""
    name = clean_name(path.name)
    stem, ext = Path(name).stem, Path(name).suffix.lower()

    m = SUBMIT_RE.match(name)
    if m:
        rest = name[m.end():-len(ext)] if ext else name[m.end():]
        if not re.search(r"[A-Za-z가-힣]", rest):  # 업로드 과정에서 한글 이름이 날아간 경우
            rest = PROFILE.get("name", "name")
        canon = f"{'prelab' if m.group(1) else 'lab'}{int(m.group(2)):02d}_{m.group(3)}_{rest}{ext}"
        return f"{int(m.group(2)):02d}", "submitted", canon, "제출 파일명"

    lab, why = (forced_lab, "--lab") if forced_lab else lab_from_name(stem)
    if lab is None:
        lab, why = lab_from_content(path)
    if lab is None and croot is not None:
        lab, why = lab_from_date(path, croot, head_text(path))
    if lab is None:
        return None, None, name, ""

    if ext in DATA_EXT:
        return lab, "report/raw", name if re.search(r"[A-Za-z0-9가-힣]", Path(name).stem) else f"data{ext}", why
    if ext in IMG_EXT:
        return lab, "report/photos", name, why

    low = stem.lower()
    if "guide" in low and "book" in low:
        new = "guidebook"
    elif "guideline" in low:
        new = "report_guideline"
    elif "introduction" in low:
        new = "introduction"
    elif re.search(r"tool", low):
        new = "experimental_tools"
    elif re.search(r"가이드\s*북|실험\s*(?:교재|지침서)|guide", low):
        new = "guidebook"
    else:
        rest = slug(stem)
        if not rest or rest == "file":
            new = "slides"
        elif DATE6.fullmatch(rest):
            new = f"slides_{rest}"
        else:
            new = rest
    return lab, "materials", new + ext, why


def pdf_text(p: Path):
    if p.suffix.lower() != ".pdf":
        return
    if find_tool("pdftotext"):
        subprocess.run([find_tool("pdftotext"), "-layout", str(p), str(p.with_suffix(".txt"))], check=False)
        return
    try:   # pdftotext가 없으면 (Windows 등) pymupdf로
        import pymupdf
        with pymupdf.open(str(p)) as doc:
            p.with_suffix(".txt").write_text("\n\f".join(pg.get_text() for pg in doc), encoding="utf-8")
    except Exception as e:
        print(f"  (텍스트 추출 실패: {p.name}: {e})")


def dest_dir(croot: Path, lab: str, sub: str, unit: str = "lab") -> Path:
    """croot = courses/<과목>, unit = 과목 course.yaml의 unit (lab, hw)"""
    if lab == "00":
        return croot / ("materials" if sub == "materials" else sub)
    return croot / f"{unit}{lab}" / sub


# ───────────────────────── requirement / meta drafts ─────────────────────────
def _date_from_text(text: str, year: int):
    m = re.search(r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{1,2}),?\s*(\d{4})?", text)
    if not m:
        return None
    y = int(m.group(3)) if m.group(3) else year
    return dt.date(y, MONTHS[m.group(1).lower()], int(m.group(2)))


def _deadline(text: str, key: str, year: int):
    """'Submission deadline: 23:59 September 27 (Sun)' 류 파싱"""
    pat = {
        "prelab": r"(?is)Prelab report.*?(?:deadline|by)[:\s]*23:59\s+([A-Za-z]+)\s+(\d{1,2})",
        "report": r"(?is)(?<!pre)Lab Report\s.*?by\s+23:59\s+([A-Za-z]+)\s+(\d{1,2})",
    }[key]
    m = re.search(pat, text)
    if not m or m.group(1).lower() not in MONTHS:
        return None
    return dt.date(year, MONTHS[m.group(1).lower()], int(m.group(2)))


def _sunday_before(d: dt.date):
    return d - dt.timedelta(days=(d.weekday() + 1) % 7 or 7)


def _sunday_after(d: dt.date):
    return d + dt.timedelta(days=(6 - d.weekday()) or 7)


def draft_meta_and_requirements(lab_dir: Path, lab: str):
    mat = lab_dir / "materials"
    guide = mat / "guidebook.txt"
    slides = sorted(mat.glob("slides*.txt"))
    gtext = guide.read_text(encoding="utf-8", errors="ignore") if guide.exists() else ""
    stext = "\n".join(s.read_text(encoding="utf-8", errors="ignore") for s in slides)

    # meta.yaml
    meta_p = lab_dir / "meta.yaml"
    if not meta_p.exists():
        year = dt.date.today().year
        title = ""
        m = re.search(r"(?im)^\s*(?:Lab|실험|Experiment)\s*\d+\s*[.:]\s*(.+?)\s*$\n?(.*)$", gtext)
        if m:
            title = (m.group(1) + (" " + m.group(2).strip() if m.group(2).strip() and not re.match(r"\d", m.group(2).strip()) else "")).strip()
        if not title:
            intro = lab_dir.parent / "materials" / "introduction.txt"
            if intro.exists():
                mm = re.search(rf"(?:Lab|실험)\s*0?{int(lab)}\s*[:.]\s*(.+)", intro.read_text(encoding="utf-8", errors="ignore"))
                title = mm.group(1).strip() if mm else ""
        lab_date = _date_from_text(stext, year)
        if lab_date is None:
            km = KO_DATE.search(stext)
            if km:
                try:
                    lab_date = dt.date(int(km.group(1) or year), int(km.group(2)), int(km.group(3)))
                except ValueError:
                    lab_date = None
        notes = []
        pre = _deadline(stext, "prelab", year)
        rep = _deadline(stext, "report", year)
        if lab_date:
            rule_pre, rule_rep = _sunday_before(lab_date), _sunday_after(lab_date)
            if pre != rule_pre:
                notes.append(f"prelab 마감: 슬라이드 {pre} / 규칙상 {rule_pre} — 확인 필요")
                pre = rule_pre if pre is None or pre > lab_date else pre
            if rep != rule_rep:
                notes.append(f"report 마감: 슬라이드 {rep} / 규칙상 {rule_rep} — 확인 필요")
                rep = rule_rep if rep is None or rep < lab_date else rep
        meta = {
            "lab": int(lab),
            "title_en": title,
            "title": title if re.search(r"[가-힣]", title) else "[확인 필요: 한국어 제목]",   # 한국어 교재면 그대로
            "lab_date": str(lab_date) if lab_date else "[확인 필요]",
            "prelab_due": f"{pre} 23:59" if pre else "[확인 필요]",
            "report_due": f"{rep} 23:59" if rep else "[확인 필요]",
        }
        if notes:
            meta["_notes"] = notes
        meta_p.write_text(yaml.safe_dump(meta, allow_unicode=True, sort_keys=False), encoding="utf-8")
        print("  + meta.yaml 초안", *[f"\n    ⚠ {n}" for n in notes])

    # requirements.md
    req_p = lab_dir / "requirements.md"
    if not req_p.exists() and (gtext or stext):
        parts = [f"# Lab {lab} 요구사항 (자동 추출 초안 — 원문과 대조 후 사용)\n"]
        for label, pat in (("Prelab", r"(?:Pre-?lab|예비\s*(?:실험|보고서|과제)|실험\s*전\s*과제|모의\s*실험\s*보고서)"),
                           ("실험 보고서 (결과보고서 문항)", r"(?:실험\s*보고서|결과\s*보고서)"),
                           ("Lab (실험)", r"(?:Lab|실험\s*(?:방법|절차|내용)?)"),
                           ("Discussion and Matters to Consider",
                            r"(?:Discussion(?: and Matters to Consider)?|결과\s*및\s*토의|토의|고찰)")):
            m = re.search(rf"(?ms)^\s*(\d+(?:\.\d+)*)\.?\s+{pat}\s*$(.*?)(?=^\s*\d+(?:\.\d+)*\.?\s+(?:[A-Z][a-z]|[가-힣])|\Z)", gtext)
            if m:
                body = re.sub(r"\n{3,}", "\n\n", m.group(2)).rstrip()
                parts.append(f"## guidebook {m.group(1)}. {label}\n\n```\n{body}\n```\n")
        keep = [ln.strip() for ln in stext.splitlines()
                if re.search(r"(?i)goal|only|must|include|answer|deadline|submit|checklist|measure|don.t have to|\d\.\d|\b[0-9]\.[a-g]\b|\b\d[a-g]\b"
                             r"|목표|제출|마감|필수|포함|답하|범위|측정|제외|하지 않아도|예비|결과\s*보고서", ln)]
        if keep:
            parts.append("## 슬라이드에서 뽑은 지시 (범위·마감·필수 항목)\n\n```\n" + "\n".join(keep) + "\n```\n")
        parts.append("## 보고서 대응표 (Claude가 채움)\n\n| 요구 항목 | 보고서 위치 | 상태 |\n|:--|:--|:--|\n")
        req_p.write_text("\n".join(parts), encoding="utf-8")
        print("  + requirements.md 초안")


def draft_hw(hw_dir: Path, num: str):
    """과제 meta.yaml과 requirements.md 초안.
    실습 자료의 'Homework' 절에서 문제(1. 2. 3.)와 제약('사용 금지', '사용하지 않고'), '제출 기한'을 뽑는다."""
    text = "\n".join(t.read_text(encoding="utf-8", errors="ignore") for t in sorted((hw_dir / "materials").glob("*.txt")))
    due = None
    m = re.search(r"제출\s*기한\s*[:：]?\s*([^\n]+)", text)
    if m:
        due = m.group(1).strip()
    meta_p = hw_dir / "meta.yaml"
    if not meta_p.exists():
        meta_p.write_text(yaml.safe_dump({"hw": int(num), "title": "[확인 필요: 과제 주제]", "due": due or "[확인 필요]"},
                                         allow_unicode=True, sort_keys=False), encoding="utf-8")
        print("  + meta.yaml 초안" + (f" (제출 기한 {due})" if due else ""))
    req_p = hw_dir / "requirements.md"
    if req_p.exists() or not text:
        return
    # 문제: 'Homework' 뒤의 '1. …' 또는 'Problem 1', '문제 1'
    hw_part = text[text.find("Homework"):] if "Homework" in text else text
    probs: dict[str, str] = {}
    for mm in re.finditer(r"(?ms)^\s*(?:(?:Problem|Prob\.?|문제)\s*)?(\d{1,2})[.)]\s+(.+?)(?=^\s*(?:(?:Problem|문제)\s*)?\d{1,2}[.)]\s|^\s*(?:■|□)?\s*eTL|\Z)", hw_part):
        body = " ".join(mm.group(2).split())
        if len(body) > 15 and mm.group(1) not in probs:
            probs[mm.group(1)] = body[:400]
    rows = []
    for n, body in probs.items():
        limits = "; ".join(re.findall(r"[^.,()]*(?:사용\s*금지|사용하지\s*않고|사용할\s*것|이상의)[^.,()]*", body))
        rows.append(f"| Problem {n} | {body[:120]}{'…' if len(body) > 120 else ''} | {limits or '-'} | 코드 `%% Problem {n}` + 보고서 절 | |")
    req_p.write_text(f"# HW{int(num)} 요구사항 (자동 추출 초안 — 과제 원문과 대조)\n\n"
                     f"제출 기한: {due or '[확인 필요]'}\n\n"
                     "| 문제 | 내용 | 제약 (지켰는지 `mcode.py check --forbid`로 확인) | 코드·보고서 위치 | 상태 |\n|:--|:--|:--|:--|:--|\n"
                     + ("\n".join(rows) or "| [확인 필요: 과제 PDF에서 문제] | | | | |") + "\n", encoding="utf-8")
    print(f"  + requirements.md 초안 (문제 {len(probs)}개)")


# ───────────────────────── main ─────────────────────────
LAB_COURSES_SKIP: set[str] = set()   # Lab·HW 번호로 분류하지 않는 과목


def process(root: Path, course: str, files: list[Path], from_inbox: bool, forced: str | None, dry: bool):
    croot = root / "courses" / course
    inbox = root / "inbox" / course
    if course not in known_courses(root):
        sys.exit(f"✗ 과목 '{course}'을 모름 — 가능한 과목: {', '.join(known_courses(root))}")
    if not dry:
        croot.mkdir(parents=True, exist_ok=True)
    reg_p = croot / ".ingest.json"
    reg = json.loads(reg_p.read_text(encoding="utf-8")) if reg_p.exists() else {}
    touched_labs = set()
    unit = str(course_defaults(course).get("unit", "lab"))   # 작업 단위 폴더: lab01/ 또는 hw01/
    seen = {}
    print(f"[{course}]")
    for f in files:
        h = sha(f)
        lab, sub, name, why = classify_hw(f, forced) if unit == "hw" else classify(f, forced, croot)
        if h in seen:
            print(f"= 중복 건너뜀: {f.name} (같은 내용: {seen[h]})")
            if from_inbox and not dry:
                f.unlink()
            continue
        seen[h] = f.name
        if h in reg:
            print(f"= 중복 건너뜀: {f.name} (이미 {reg[h]['path']})")
            if from_inbox and not dry:
                f.unlink()
            continue
        if lab is None:
            print(f"? Lab 번호 못 찾음: {f.name} → inbox/{course}/_unsorted\n"
                  f"    에이전트: 파일을 열어 주제·날짜를 course.md 일정과 대조해 Lab을 정하고 "
                  f"`ingest.py --course {course} --lab NN inbox/{course}/_unsorted/{f.name}`로 다시 (사용자에게 묻지 않는다)")
            if not dry and f.parent.name != "_unsorted":
                (inbox / "_unsorted").mkdir(parents=True, exist_ok=True)
                shutil.move(str(f), inbox / "_unsorted" / f.name) if from_inbox else shutil.copy2(f, inbox / "_unsorted" / f.name)
            continue
        d = dest_dir(croot, lab, sub, unit)
        target = d / name
        if target.exists() and sha(target) != h:  # 같은 이름 다른 내용 → 버전 보존
            target = d / f"{target.stem}_{dt.date.today():%y%m%d}{target.suffix}"
        print(f"→ {f.name}  ⇒  {target.relative_to(root)}" + (f"   [{why}]" if why and not why.startswith("파일명 '[") else ""))
        if dry:
            continue
        d.mkdir(parents=True, exist_ok=True)
        (shutil.move if from_inbox or f.parent.name == "_unsorted" else shutil.copy2)(str(f), target)   # _unsorted에서 다시 돌리면 옮긴다
        pdf_text(target)
        reg[h] = {"path": str(target.relative_to(root)), "original": f.name, "at": dt.datetime.now().isoformat(timespec="seconds")}
        if lab != "00":
            touched_labs.add(lab)

    if dry:
        return
    reg_p.write_text(json.dumps(reg, ensure_ascii=False, indent=1), encoding="utf-8")

    for lab in sorted(touched_labs):
        ld = croot / f"{unit}{lab}"
        subs = ("code", "figs", "materials", "submitted") if unit == "hw" else \
            ("prelab/figs", "report/figs", "report/photos", "report/scope", "report/raw", "submitted")
        for sub in subs:
            (ld / sub).mkdir(parents=True, exist_ok=True)
        print(f"  [{unit}{lab}]")
        (draft_hw if unit == "hw" else draft_meta_and_requirements)(ld, lab)

    # 자료 목록
    for idx_dir in [croot / "materials", *sorted(croot.glob(f"{unit}*/materials"))]:
        if idx_dir.exists():
            rows = [f"| {p.name} | {p.stat().st_size // 1024} KB |" for p in sorted(idx_dir.iterdir())
                    if p.is_file() and p.suffix != ".txt" and p.name != "INDEX.md"]
            (idx_dir / "INDEX.md").write_text("| 파일 | 크기 |\n|:--|--:|\n" + "\n".join(rows) + "\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="*", type=Path)
    ap.add_argument("--course", help="과목 폴더 이름 (logic, circuit …). 생략하면 inbox/의 모든 과목")
    ap.add_argument("--lab", help="Lab 번호 강제 지정 (예: 01)")
    ap.add_argument("--root", type=Path)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    root = (a.root or find_root(Path.cwd())).resolve()
    remember_engine(find_workspace(root))
    if (root / "profile.yaml").exists():
        PROFILE.update(yaml.safe_load((root / "profile.yaml").read_text(encoding="utf-8")) or {})
    forced = f"{int(a.lab):02d}" if a.lab else None

    if a.files:
        if not a.course:
            sys.exit("✗ 파일을 직접 줄 때는 --course를 같이 준다 (예: --course logic)")
        process(root, a.course, a.files, False, forced, a.dry_run)
        return

    inbox_root = root / "inbox"
    courses = [a.course] if a.course else sorted(d.name for d in inbox_root.iterdir()
                                                if d.is_dir() and d.name not in LAB_COURSES_SKIP) if inbox_root.exists() else []
    done = False
    for c in courses:
        inbox = root / "inbox" / c
        files = sorted(p for p in inbox.glob("*") if p.is_file() and not p.name.startswith("."))
        if files:
            process(root, c, files, True, forced, a.dry_run)
            done = True
    if not done:
        print("처리할 파일 없음 (inbox/<과목>/)")


if __name__ == "__main__":
    sys.exit(main())
