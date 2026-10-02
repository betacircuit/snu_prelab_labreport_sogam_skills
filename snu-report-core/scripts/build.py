#!/usr/bin/env python3
"""마크다운 원고 → 스타일 적용된 docx → pdf

사용:
  python build.py <lab_dir> prelab           # <lab_dir>/prelab/prelab.md  → build/prelab01_학번_이름.docx
  python build.py <lab_dir> report           # <lab_dir>/report/report.md  → build/lab01_학번_이름.docx
  python build.py <lab_dir> report --pdf     # + 미리보기 PDF (build/preview/) — 확인용
  python build.py <lab_dir> report --final   # 미해결 상태 표시가 남아 있으면 실패

사용자에게 주는 결과물은 항상 .docx (수정 가능). PDF 변환·제출은 사용자가 Word에서 한다.

원고 문법 (pandoc markdown + 아래 확장):
  # 제목               → 자동 번호 "1." / "1.1)" (style.yaml numbering.heading)
  # 4.2 3-bit comparator → "2. 3-bit comparator" (원문 번호는 원고·requirements.yaml에 보존)
  ## 3) 120 nF           → "2.3) 120 nF" (제목 번호는 한 번만 표시)
  # 참고문헌 {-}        → 번호 없는 제목
  ![캡션](figs/a.png){#fig:and width=60%}   → 틀 안 "Fig.1 - 캡션" (style.yaml caption_format)
  Table: 캡션 {#tbl:tt}                      → 표 맨 아래 캡션 행 "Table.1 - 캡션"
  @fig:and, @tbl:tt                          → "Fig.1", "Table.1"
  $\\overline{A}B + A\\overline{B}$            → Word 수식
  상태 표시 (references/evidence.md):
    [TODO: …] [미제공: …] [확인 필요: …] [판독 불가: …]  → 초안에서 형광펜, --final 이면 빌드 중단
    [미측정: …] [미수행: …]                            → 초안에서 형광펜, --final 이면 경고
                                                          (최종본에서는 문장으로 풀어 쓸 것)
"""
from __future__ import annotations

import _console  # noqa: F401  (Windows에서 한글·기호 출력)
import argparse
import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
TEMPLATES = HERE.parent / "templates"
sys.path.insert(0, str(HERE))
from docx_post import break_before_frames, postprocess, split_captions  # noqa: E402
from make_template import build_reference  # noqa: E402
from ws import course_defaults, find_tool, find_workspace, remember_engine  # noqa: E402

KIND_LABEL = {"prelab": "Prelab Report", "report": "Lab Report", "hw": "Homework"}
KIND_KO = {"prelab": "예비보고서", "report": "결과보고서", "hw": "결과 보고서"}
SRC = {"prelab": "prelab/prelab.md", "report": "report/report.md", "hw": "hw.md"}
BLOCKING = ("TODO", "미제공", "확인 필요", "판독 불가")
SOFT = ("미측정", "미수행")
MARK_RE = re.compile(r"\[(" + "|".join(BLOCKING + SOFT) + r")(?::[^\]]*)?\]")


# ───────────────────────── config ─────────────────────────
def find_upward(start: Path, name: str) -> Path | None:
    for d in [start, *start.parents]:
        if (d / name).exists():
            return d / name
    return None


class _D(dict):
    def __missing__(self, k):
        return ""


def load_vars(lab_dir: Path, kind: str) -> dict:
    meta = yaml.safe_load((lab_dir / "meta.yaml").read_text(encoding="utf-8")) or {}
    prof_path = find_upward(lab_dir, "profile.yaml")
    prof = yaml.safe_load(prof_path.read_text(encoding="utf-8")) if prof_path else {}
    # 과목 설정: 과목 스킬의 course.yaml(과목명, 파일명 규칙) ← 작업 폴더 courses/<과목>/course.yaml(조, 조원)
    course_path = find_upward(lab_dir, "course.yaml")
    over = (yaml.safe_load(course_path.read_text(encoding="utf-8")) or {}) if course_path else {}
    key = str(over.get("key") or lab_dir.resolve().parent.name)
    course = {**course_defaults(key), **over, "key": key}
    v = _D({**(prof or {}), **course, **meta})
    v["kind"] = KIND_LABEL[kind]
    v["kind_key"] = kind
    lab = meta.get("lab", meta.get("hw", ""))
    v["lab"] = f"{int(lab):02d}" if str(lab).isdigit() else str(lab)
    mates = v.get("teammates") or []
    v["teammates"] = ", ".join(m["name"] if isinstance(m, dict) else str(m) for m in mates)
    v["lab_date"] = str(v.get("lab_date", ""))
    v["lab_num"] = str(int(lab)) if str(lab).isdigit() else str(lab)
    v["kind_ko"] = KIND_KO[kind]
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", v["lab_date"])
    v["lab_date_ko"] = f"{m.group(1)}년 {int(m.group(2))}월 {int(m.group(3))}일" if m else v["lab_date"]
    v["team"] = str(v.get("team", "") or "")
    return v


def fill_line(tmpl: str, v: dict) -> str | None:
    """변수가 하나라도 비어 있으면 None (그 줄 생략)."""
    keys = re.findall(r"\{(\w+)\}", tmpl)
    if any(not str(v.get(k, "")).strip() or str(v.get(k, "")).startswith("[") for k in keys):
        return None
    return tmpl.format_map(v)


# ───────────────────────── 제출 파일명 (Lab00 가이드라인: 틀리면 감점) ─────────────────────────
FILENAME_RE = re.compile(r"^(prelab|lab)\d{2}_\d{4}-\d{5}_[가-힣A-Za-z]+$")
HW_FILENAME_RE = re.compile(r"^HW\d+_[가-힣A-Za-z]+_\d{4}-\d{5}$")   # 기초전자기학: HW1_홍길동_2025-12345


DEFAULT_FILENAME = {"prelab": "prelab{lab}_{student_id}_{name}", "report": "lab{lab}_{student_id}_{name}",
                    "hw": "HW{lab_num}_{name}_{student_id}"}


def submission_stem(kind: str, v: dict) -> str:
    """course.yaml의 filename 규칙. prelab/report는 prelab01_2025-12345_홍길동 형식을 검사한다."""
    tmpl = (v.get("filename") or {}).get(kind) or DEFAULT_FILENAME[kind]
    stem = tmpl.format_map(v)
    if kind in ("prelab", "report") and not FILENAME_RE.match(stem):
        sys.exit(f"✗ 제출 파일명 규칙 위반: {stem!r} — profile.yaml의 student_id(2025-12345 형식)/name, course.yaml의 filename, meta.yaml의 lab 확인")
    if kind == "hw" and not HW_FILENAME_RE.match(stem):
        sys.exit(f"✗ 제출 파일명 규칙 위반: {stem!r} — HW1_홍길동_2025-12345 꼴 (profile.yaml, meta.yaml의 hw 확인)")
    return stem



# ───────────────────────── 수식 띄어쓰기 ─────────────────────────
_BINOPS = [r"\approx", r"\times", r"\oplus", r"\cdot", r"\neq", r"\leq", r"\geq", r"\le", r"\ge", "=", "+", "<", ">"]
MATH_RE = re.compile(r"(?<!\\)(\$\$.+?\$\$|\$[^$\n]+?\$)", re.S)


def space_math(expr: str) -> str:
    """이항 연산자(=, +, −, ×, ≈ …) 양옆과 단위 앞에 공백. 첨자 안은 건드리지 않는다."""
    sp = r"\text{ }"   # Word·LibreOffice·한컴 모두에서 보이는 공백 (\; 같은 수식 공백은 일부 뷰어가 무시)
    expr = expr.replace("\\qquad", r"\text{      }").replace("\\quad", r"\text{   }")
    expr = expr.replace("\\ ", sp).replace("\\,", sp).replace("\\;", sp)
    out, i, n = [], 0, len(expr)
    while i < n:
        c = expr[i]
        if c in "_^" and i + 1 < n and expr[i + 1] == "{":  # 첨자 블록은 그대로
            depth, j = 0, i + 1
            while j < n:
                depth += expr[j] == "{"
                depth -= expr[j] == "}"
                j += 1
                if depth == 0:
                    break
            out.append(expr[i:j])
            i = j
            continue
        op = next((o for o in _BINOPS if expr.startswith(o, i) and not (o[0] == "\\" and i + len(o) < n and expr[i + len(o)].isalpha())), None)
        prev = "".join(out).rstrip()
        if op is None and c == "-" and prev and (prev[-1].isalnum() or prev[-1] in "})") and not prev.endswith(sp):   # 연산자 바로 뒤의 -는 부호
            op = "-"
        if op:
            while out and out[-1] in (" ", sp):
                out.pop()
            out.append(sp + op + sp)
            i += len(op)
            while i < n and expr[i] == " ":
                i += 1
            if expr.startswith(sp, i):
                i += len(sp)
            continue
        if expr.startswith(sp, i) and out and out[-1].endswith(sp):
            i += len(sp)
            continue
        out.append(c)
        i += 1
    return "".join(out)


# ───────────────────────── 수식 글자 세우기 ─────────────────────────
# 변수만 기울임. 함수 이름(sin, sinc …)과 단위(V/m, F/m, N m^2/C^2 …)는 바로 세운다 (사용자 지적: 과도한 기울임).
_FUNCS = ("sinh", "cosh", "tanh", "sin", "cos", "tan", "cot", "sec", "csc", "exp", "log", "ln", "arctan", "arcsin", "arccos")
_FUNC_RE = re.compile(r"(?<![\\A-Za-z])(sinc|" + "|".join(_FUNCS) + r")(?![A-Za-z])")
_UATOM = r"(?:[pnuµμmckMG]?(?:Hz|Wb|rad|dB|V|A|Ω|F|H|s|m|C|N|J|W|T|S|g|K)|\\Omega)(?:\^\{?-?\d+\}?)?"
_UNIT_RE = re.compile(r"(?P<pre>[\d}])(?P<sp>\s*(?:\\[ ,;:]|~)\s*|\s+)(?P<u>" + _UATOM +
                      r"(?:\s*(?:/|\\cdot|\\[ ,;])\s*" + _UATOM + r")*)(?![A-Za-z{])")
_PROTECT_RE = re.compile(r"\\(?:text|mathrm|operatorname|mathit|mathbf)\s*\{[^{}]*\}")


def upright_math(expr: str) -> str:
    """수식 안 함수 이름은 \\sin·\\operatorname{sinc}로, 숫자 뒤 단위는 \\mathrm{…}으로 (이미 \\text·\\mathrm이면 그대로)."""
    parts, last = [], 0
    for m in _PROTECT_RE.finditer(expr):
        parts.append((expr[last:m.start()], True))
        parts.append((m.group(0), False))
        last = m.end()
    parts.append((expr[last:], True))
    out = []
    for seg, edit in parts:
        if edit:
            seg = _FUNC_RE.sub(lambda m: r"\operatorname{sinc}" if m.group(1) == "sinc" else "\\" + m.group(1), seg)
            seg = _UNIT_RE.sub(lambda m: m["pre"] + r"\ " + r"\mathrm{" + re.sub(r"\\ ", r"\\,", m["u"].strip()) + "}", seg)
        out.append(seg)
    return "".join(out)


def space_all_math(text: str) -> str:
    parts = re.split(r"(```.*?```)", text, flags=re.S)
    for k in range(0, len(parts), 2):
        parts[k] = MATH_RE.sub(lambda m: (m.group(0)[:2] + space_math(upright_math(m.group(0)[2:-2])) + m.group(0)[-2:])
                               if m.group(0).startswith("$$") else "$" + space_math(upright_math(m.group(0)[1:-1])) + "$", parts[k])
    return "".join(parts)

NBSP = "\u00a0"
UNIT_RE = re.compile(r"(?<=\d) (?=(?:ns|us|ms|s|mV|V|mA|Ω|kΩ|MΩ|pF|nF|uF|Hz|kHz|MHz)(?![A-Za-z]))")
SHORT_EQ_RE = re.compile(r"(?<![A-Za-z0-9_\\{])([A-Za-z]) = ([A-Za-z0-9.]{1,4})(?![A-Za-z0-9_.])")


def nbsp_text(text: str) -> str:
    """수식 밖의 짧은 식(B = 1)과 숫자·단위(10 ns) 사이 공백을 줄바꿈 없는 공백으로."""
    parts = re.split(r"(```.*?```|\$\$.*?\$\$|\$[^$\n]+\$)", text, flags=re.S)
    for k in range(0, len(parts), 2):
        seg = UNIT_RE.sub(NBSP, parts[k])
        parts[k] = SHORT_EQ_RE.sub(lambda m: f"{m.group(1)}{NBSP}={NBSP}{m.group(2)}", seg)
    return "".join(parts)


# ───────────────────────── preprocessing ─────────────────────────
IMG_RE = re.compile(r"^(\s*)!\[(?P<cap>.*?)\]\((?P<src>[^)]+)\)(?P<attr>\{[^}]*\})?\s*$")
TBLCAP_RE = re.compile(r"^Table:\s*(?P<cap>.*?)\s*(?P<attr>\{[^}]*\})?\s*$")
HEAD_RE = re.compile(r"^(?P<hash>#{1,3})\s+(?P<text>.+?)\s*(?P<attr>\{[^}]*\})?\s*$")
ID_RE = re.compile(r"#((?:fig|tbl):[A-Za-z0-9_-]+)")


def _strip_id(attr: str | None) -> str:
    if not attr:
        return ""
    rest = ID_RE.sub("", attr[1:-1]).strip()
    return "{" + rest + "}" if rest else ""


def preprocess(md: str, st: dict) -> tuple[str, list[str]]:
    from proof import strip_heading_labels
    num = st["numbering"]
    fig_p, tbl_p = num["figure_prefix"], num["table_prefix"]
    cap_fmt = num.get("caption_format", "{prefix}.{n} - {text}")
    ref_fmt = num.get("ref_format", "{prefix}.{n}")
    labels: dict[str, str] = {}
    out, warnings = [], []
    nfig = ntbl = 0
    h = [0, 0, 0]
    in_code = False

    for line in md.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
            out.append(line)
            continue
        if in_code:
            out.append(line)
            continue

        m = HEAD_RE.match(line)
        if m:
            attr = m["attr"] or ""
            inner = attr[1:-1] if attr else ""
            unnumbered = bool(re.search(r"(^|\s)-(\s|$)", inner)) or ".unnumbered" in inner
            lvl = len(m["hash"])
            if num["heading"] and not unnumbered:
                h[lvl - 1] += 1
                for i in range(lvl, 3):
                    h[i] = 0
                fmts = num.get("heading_formats", ["{0}.", "{0}.{1}", "{0}.{1}.{2}"])
                n = fmts[lvl - 1].format(*h[:lvl])
                line = f"{m['hash']} {n} {strip_heading_labels(m['text'])} {attr}".rstrip()
            out.append(line)
            continue

        m = IMG_RE.match(line)
        if m and m["cap"]:
            nfig += 1
            for i in ID_RE.findall(m["attr"] or ""):
                if i in labels:
                    warnings.append(f"중복 라벨: {i}")
                labels[i] = ref_fmt.format(prefix=fig_p, n=nfig)
            cap = cap_fmt.format(prefix=fig_p, n=nfig, text=m["cap"])
            if m["src"].startswith("snu-missing:"):
                # 초안의 누락 그림도 참조 번호는 유지하되 가짜 그림/빈 틀을 만들지 않는다.
                out.append('::: {custom-style="Missing Figure"}\n' +
                           f'[⟦미제공: {cap}⟧]{{custom-style="Marker"}}\n:::')
            else:
                out.append(f"{m.group(1)}![{cap}]({m['src']}){_strip_id(m['attr'])}")
            continue

        m = TBLCAP_RE.match(line)
        if m:
            ntbl += 1
            for i in ID_RE.findall(m["attr"] or ""):
                if i in labels:
                    warnings.append(f"중복 라벨: {i}")
                labels[i] = ref_fmt.format(prefix=tbl_p, n=ntbl)
            out.append("Table: " + cap_fmt.format(prefix=tbl_p, n=ntbl, text=m["cap"]))
            continue

        # 상태 표시 형광펜
        line = MARK_RE.sub(lambda mm: f'[{mm.group(0).replace("[", "⟦").replace("]", "⟧")}]{{custom-style="Marker"}}', line)
        out.append(line)

    text = "\n".join(out)

    def ref(mm):
        key = mm.group(1)
        if key not in labels:
            warnings.append(f"정의되지 않은 참조: @{key}")
            return f"??{key}"
        return labels[key]

    text = re.sub(r"@((?:fig|tbl):[A-Za-z0-9_-]+)", ref, text)
    if st.get("math", {}).get("op_space", True):
        text = space_all_math(text)
    text = nbsp_text(text)
    return text, warnings


def title_md(v: dict, st: dict) -> str:
    tb = st.get("title_block", {})
    kind = v["kind_key"]
    title = tb.get(f"{kind}_title", "Lab {lab} : {title}").format_map(v)
    def pick(t):  # 한 줄에 후보 여러 개면 값이 다 채워지는 첫 후보 (조원이 없으면 조원 없는 줄로)
        for c in (t if isinstance(t, list) else [t]):
            got = fill_line(c, v)
            if got:
                return got
        return None
    lines = [x for x in (pick(t) for t in tb.get(f"{kind}_lines", [])) if x]
    if tb.get("mode", "block") == "page":
        meta = "\n".join(f"{x}\\" for x in lines).rstrip("\\")
        pb = '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```'
        return (f'::: {{custom-style="Cover Kicker"}}\n{v.get("course", "")} · {v["kind"]}\n:::\n\n'
                f'::: {{custom-style="Title"}}\n{title}\n:::\n\n&nbsp;\n\n'
                f'::: {{custom-style="Cover Meta"}}\n{meta}\n:::\n\n{pb}\n')
    def gap(t: str) -> str:  # 두 칸 이상 공백은 em space로 (마크다운이 공백을 합치지 않게)
        return re.sub(r" {2,}", lambda mm: "\u2003" * (len(mm.group(0)) // 2), t)

    title_lines = "\\\n".join(gap(x.strip()) for x in title.split("\n") if x.strip())
    out = f'::: {{custom-style="Doc Title"}}\n{title_lines}\n:::\n\n'
    sub_t = tb.get(f"{kind}_subtitle")
    if sub_t:
        sub = fill_line(sub_t, v)
        if sub:
            out += f'::: {{custom-style="Doc Subtitle"}}\n{gap(sub)}\n:::\n\n'
    meta = "\n".join(f"{gap(x)}\\" for x in lines).rstrip("\\")
    if meta:
        out += f'::: {{custom-style="Doc Meta"}}\n{meta}\n:::\n\n'
    if tb.get("rule", False):
        out += '::: {custom-style="Doc Rule"}\n&nbsp;\n:::\n\n'
    return out


# ───────────────────────── postprocess ─────────────────────────
def one_column_cover(docx_path: Path):
    """본문이 다단일 때 표지는 1단으로 유지 (페이지 나누기 → 구역 나누기로 교체)."""
    doc = Document(docx_path)
    body_sect = doc.sections[-1]._sectPr
    for p in doc.paragraphs:
        brs = p._p.findall(".//" + qn("w:br"))
        if any(b.get(qn("w:type")) == "page" for b in brs):
            for r in list(p._p.findall(qn("w:r"))):
                p._p.remove(r)
            sp = copy.deepcopy(body_sect)
            cols = sp.find(qn("w:cols"))
            if cols is not None:
                cols.set(qn("w:num"), "1")
            p._p.get_or_add_pPr().append(sp)
            tp = body_sect.find(qn("w:titlePg"))
            if tp is not None:
                body_sect.remove(tp)
            break
    doc.save(docx_path)


def unescape_markers(docx_path: Path):
    """형광펜 표시용으로 바꿔 둔 ⟦ ⟧ 를 원래 대괄호로 되돌림."""
    doc = Document(docx_path)
    for t in doc.element.body.iter(qn("w:t")):
        if t.text and ("⟦" in t.text or "⟧" in t.text):
            t.text = t.text.replace("⟦", "[").replace("⟧", "]")
    doc.save(docx_path)


def to_pdf(docx_path: Path, output_dir: Path | None = None, soffice: str | None = None) -> Path:
    output_dir = output_dir or docx_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    tool = soffice or find_tool("soffice")
    if not tool:
        raise RuntimeError("LibreOffice가 없음 — 호스트의 렌더러 또는 SNU_SOFFICE 경로를 지정")
    result = output_dir / (docx_path.stem + ".pdf")
    result.unlink(missing_ok=True)  # 이전 실행의 PDF를 성공으로 오인하지 않는다.
    with tempfile.TemporaryDirectory(prefix="snu-lo-") as profile:
        converted = subprocess.run(
            [tool, "-env:UserInstallation=" + Path(profile).as_uri(), "--headless",
             "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path.resolve())],
            capture_output=True, text=True, errors="replace", timeout=120,
        )
    if converted.returncode or not result.exists() or result.stat().st_size == 0:
        raise RuntimeError("PDF 변환 실패: " + (converted.stdout + converted.stderr)[-1500:])
    return result


# ───────────────────────── 초안의 누락 그림 표시 ─────────────────────────
def photo_placeholders(md: str, lab_dir: Path, src_dir: Path, st: dict) -> tuple[str, list[str]]:
    """없는 그림은 참조 번호가 유지되는 한 줄 표시로, 촬영 지시는 CLI 목록으로 보낸다."""
    out, missing = [], []
    for line in md.splitlines():
        m = IMG_RE.match(line)
        if m:
            src = m["src"].strip()
            if not any((d / src).is_file() for d in (src_dir, lab_dir)):
                hint_m = re.search(r'hint="([^"]*)"', m["attr"] or "")
                print("  사진 필요:", m["cap"], "—", hint_m.group(1) if hint_m else src, file=sys.stderr)
                attr = re.sub(r'\s*hint="[^"]*"', "", m["attr"] or "")
                line = f"{m.group(1)}![{m['cap']}](snu-missing:{Path(src).stem}){attr}"
                missing.append(src)
        out.append(line)
    return "\n".join(out), missing


# ───────────────────────── main ─────────────────────────
def build(lab_dir: Path, kind: str, final=False, pdf=False, style_path: Path | None = None) -> Path:
    lab_dir = lab_dir.resolve()
    src = lab_dir / SRC[kind]
    src_dir = src.parent
    style_path = style_path or TEMPLATES / "style.yaml"
    st = yaml.safe_load(style_path.read_text(encoding="utf-8"))
    if kind == "hw":   # 과제는 'Problem 1)'이 곧 절 번호라 자동 번호(1.)를 붙이지 않는다
        st["numbering"]["heading"] = False
    v = load_vars(lab_dir, kind)
    if "heading_numbering" in v:
        st["numbering"]["heading"] = bool(v["heading_numbering"])

    raw = src.read_text(encoding="utf-8")
    from quality import check_requirements
    requirement_errors = check_requirements(lab_dir, kind, raw)
    for error in requirement_errors:
        print("  문항 대응 ⚠", error, file=sys.stderr)
    if final and requirement_errors:
        sys.exit("✗ --final: 원문 범위·문항 대응·근거를 먼저 확정한다")
    code_free = re.sub(r"(?s)```.*?```", "", raw)
    marks = MARK_RE.findall(code_free)
    blocking = [m for m in MARK_RE.finditer(code_free) if m.group(1) in BLOCKING]
    soft = [m for m in MARK_RE.finditer(code_free) if m.group(1) in SOFT]
    if marks:
        print(f"⚠ 상태 표시 {len(marks)}개 (차단 {len(blocking)}, 미측정/미수행 {len(soft)})", file=sys.stderr)
        for m in (blocking + soft)[:15]:
            print("   ", m.group(0), file=sys.stderr)
    if final and blocking:
        sys.exit("✗ --final: [TODO]/[미제공]/[확인 필요]/[판독 불가]를 모두 해결한 뒤 다시 빌드하세요.")
    if final and soft:
        print("⚠ --final: [미측정]/[미수행] 표시는 최종본에서 문장으로 풀어 쓰는 것을 권장", file=sys.stderr)

    from style_check import check as style_check  # 문체 검사 (writing.md)
    for w in style_check(raw):
        print("  문체 ⚠", w, file=sys.stderr)
    from proof import check_text, docx_text, normalize_headings  # 오탈자 검사 (proofreading.md)
    raw, head_fixes = normalize_headings(raw)   # 가이드북 문항 번호: '4.2 …' → '4.2) …'
    for c in head_fixes:
        print("  제목 번호 고침:", c, file=sys.stderr)
    proof_warns = check_text(raw)
    for w in proof_warns:
        if not w.startswith("[제목 번호]"):
            print("  오탈자 ⚠", w, file=sys.stderr)
    meta_talk = [w for w in proof_warns if w.startswith("[본인 글]")]
    if meta_talk:   # 보고서는 본인이 쓴 글이어야 한다 — '사용자', AI, '수정 전 코드' 같은 말이 있으면 만들지 않는다
        sys.exit(f"✗ 본인이 쓴 글로 읽히지 않는 표현 {len(meta_talk)}곳 — 위 [본인 글]을 고친 뒤 다시 빌드")

    code = lab_dir / "code" / f"HW{v['lab_num']}.m" if kind == "hw" else None
    if code is not None:   # 기초전자기학 HW: 같이 내는 MATLAB 코드 검사 (mcode.py)
        from mcode import check as mcheck
        if not code.exists():
            if final:
                sys.exit(f"✗ --final: 제출할 본인 코드 파일이 없다: {code}")
            print(f"⚠ 코드 파일이 없다: {code}", file=sys.stderr)
        else:
            notes, info = mcheck(code)
            for w in notes:
                print("  코드 — 본인에게 알릴 것 (코드는 고치지 않는다):", w, file=sys.stderr)
            for w in info:
                print("  코드 — 참고:", w, file=sys.stderr)
        # 기전연 보고서의 코드는 MATLAB 화면 사진으로만 넣는다 (사용자 지정). 글자 코드는 meta.yaml에 code_text: true일 때만
        text_code = re.search(r"\{\{\s*code\s*:|^```", raw, re.M)
        if text_code and not v.get("code_text"):
            sys.exit("✗ 기전연 보고서에 글자 코드({{code: …}} 또는 ``` 블록)가 있다 — 코드는 MATLAB 화면 사진(figs/p1_code.png)으로 넣는다")
        # 문제마다 본인이 올린 코드 사진(figs/pN_code…)과 결과 사진(figs/pN_result…)이 있어야 한다
        missing_photos = []
        for m in re.finditer(r"(?ms)^#\s+Problem\s+(\d+)\b(.*?)(?=^#\s|\Z)", raw):
            n, sec = m.group(1), m.group(2)
            for kind_, label in (("code", "코드 화면"), ("result", "실행 결과")):
                refs = re.findall(rf"\]\((figs/p{n}_{kind_}[^)\s]*)\)", sec)
                if not refs:
                    missing_photos.append(f"문제 {n} {label} 사진 (figs/p{n}_{kind_}.png)")
                missing_photos += [f"문제 {n} {label} 사진 파일 {r} 없음" for r in refs if not (lab_dir / r).exists()]
        if missing_photos:
            msg = "본인이 찍은 MATLAB 사진이 아직 없다 — 대신 그리지 말고 본인에게 받는다:\n  - " + "\n  - ".join(missing_photos)
            if final:
                sys.exit("✗ --final: " + msg)
            print("⚠ " + msg + "\n  (초안에는 한 줄 누락 표시를 넣는다)", file=sys.stderr)
        from mcode import include_code   # code_text: true일 때만 쓰인다
        raw, missing = include_code(raw, code)
        for k in missing:
            print(f"⚠ {code.name}에 Problem {k} 절이 없다", file=sys.stderr)

    raw, placeholders = photo_placeholders(raw, lab_dir, src_dir, st)
    if placeholders:
        print(f"⚠ 누락 그림 {len(placeholders)}개 (초안 한 줄 표시): " + ", ".join(placeholders), file=sys.stderr)
        if final:
            sys.exit("✗ --final: 사진이 아직 없는 자리가 있다 — 사진을 받아 figs/에 넣고 다시 빌드")

    body, warns = preprocess(raw, st)
    for w in warns:
        print("⚠", w, file=sys.stderr)
    if final and warns:
        sys.exit("✗ --final: 그림/표 라벨과 참조 오류를 먼저 해결한다")
    full = title_md(v, st) + "\n" + body

    stem = submission_stem(kind, v)
    out_dir = lab_dir / "build"
    out_dir.mkdir(exist_ok=True)
    out_docx = out_dir / f"{stem}.docx"

    with tempfile.TemporaryDirectory() as td:
        ref = build_reference(style_path, Path(td) / "ref.docx", v)
        md_tmp = Path(td) / "doc.md"
        md_tmp.write_text(full, encoding="utf-8")
        subprocess.run(
            [find_tool("pandoc") or "pandoc", str(md_tmp),
             "-f", "markdown+tex_math_dollars+pipe_tables+grid_tables+fenced_divs+raw_attribute+implicit_figures+bracketed_spans",
             "-o", str(out_docx), "--reference-doc", str(ref),
             "--resource-path", os.pathsep.join([str(src_dir), str(lab_dir)])],
            check=True,
        )
    unescape_markers(out_docx)
    res = postprocess(out_docx, st)
    if res.get("frames"):
        print(f"  표 {res['frames'][0]}개, 그림 {res['frames'][1]}개를 캡션 틀에 넣음")
    if st.get("title_block", {}).get("mode") == "page" and int(st["page"].get("columns", 1)) > 1:
        one_column_cover(out_docx)
    print("✓", out_docx)
    for w in check_text(docx_text(out_docx), is_docx=True):   # 빌드 결과를 다시 읽어 확인 (원고 검사와 겹치지 않는 것만)
        if w.startswith(("[제목 번호]", "[참조]", "[표시]", "[괄호]")):
            print("  결과물 ⚠", w, file=sys.stderr)
    # 렌더링과 검토가 끝나기 전 사본은 초안 경로에만 둔다.
    ws = find_workspace(lab_dir)
    if ws:
        remember_engine(ws)
        deliver = ws / "out" / "drafts" / str(v.get("key") or lab_dir.parent.name)
        deliver.mkdir(parents=True, exist_ok=True)
        shutil.copy2(out_docx, deliver / out_docx.name)
        print("  초안 사본:", (deliver / out_docx.name).relative_to(ws))
    if pdf and find_tool("soffice"):
        # 쪽 갈림을 실제 전달할 DOCX에 보정하고 그 동일 파일에서 PDF를 만든다.
        prev_dir = out_dir / "preview"
        prev_dir.mkdir(exist_ok=True)
        out_pdf = to_pdf(out_docx, output_dir=prev_dir)
        for _ in range(8 if out_pdf.exists() else 0):
            splits = split_captions(out_pdf, st)
            if not splits or not break_before_frames(out_docx, splits[:1], st):
                break
            out_pdf = to_pdf(out_docx, output_dir=prev_dir)
        print("  미리보기:", out_pdf) if out_pdf.exists() else print("⚠ PDF 변환 실패 (LibreOffice) — docx는 그대로 쓸 수 있다. Linux면 Writer가 빠졌을 수 있다: apt install libreoffice-writer", file=sys.stderr)
    if kind == "hw" and code is not None and code.exists() and ws:
        draft_code = ws / "out" / "drafts" / str(v.get("key") or lab_dir.parent.name) / code.name
        draft_code.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(code, draft_code)
        if not (out_dir / f"{stem}.zip").exists():
            print("  제출 zip은 quality.py render → 전 쪽 review → deliver 이후 생성한다")
    from quality import check_docx, record_build
    structure = check_docx(out_docx)
    if final and structure["errors"]:
        sys.exit("✗ --final: " + "; ".join(structure["errors"]))
    record_build(out_docx, lab_dir, kind, style_path)
    if ws:
        shutil.copy2(out_docx, deliver / out_docx.name)  # 배치 보정 이후 동일 초안
    if pdf and not find_tool("soffice"):
        print("⚠ 미리보기 미생성: 호스트의 렌더러로 quality.py render를 실행한다", file=sys.stderr)
    print("  완료 전: quality.py render → 모든 쪽 이미지 확인 → review → deliver")
    return out_docx


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("lab_dir", type=Path)
    ap.add_argument("kind", choices=["prelab", "report", "hw"])
    ap.add_argument("--final", action="store_true", help="미해결 상태 표시가 남아 있으면 실패")
    ap.add_argument("--pdf", action="store_true", help="미리보기 PDF도 만든다 (build/preview/)")
    ap.add_argument("--style", type=Path)
    a = ap.parse_args()
    build(a.lab_dir, a.kind, final=a.final, pdf=a.pdf, style_path=a.style)
